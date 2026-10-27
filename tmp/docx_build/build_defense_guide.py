from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path(r"E:\PrivacyPreserving-Healthcare-FL\output\documents\CAABC-FL_Thesis_Defense_Preparation_Guide_Bangla.docx")
FONT = "Nirmala UI"
MONO = "Consolas"
NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE = "EEF5FA"
GRAY = "666666"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, italic=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("পৃষ্ঠা ")
    set_run_font(run, 8.5, color=GRAY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.72)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:cs"), FONT)
normal.font.size = Pt(10.2)
normal.paragraph_format.space_after = Pt(4.5)
normal.paragraph_format.line_spacing = 1.08

for style_name, size, color, before, after in [
    ("Title", 25, NAVY, 0, 14),
    ("Subtitle", 13, BLUE, 0, 10),
    ("Heading 1", 18, NAVY, 14, 7),
    ("Heading 2", 14, BLUE, 10, 5),
    ("Heading 3", 11.5, NAVY, 8, 3),
]:
    s = styles[style_name]
    s.font.name = FONT
    s._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    s._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    s._element.rPr.rFonts.set(qn("w:cs"), FONT)
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor.from_string(color)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

footer = section.footer
fp = footer.paragraphs[0]
fp.clear()
left = fp.add_run("CAABC-FL Thesis Defense Preparation Guide")
set_run_font(left, 8.5, color=GRAY)
fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
page_p = footer.add_paragraph()
add_page_number(page_p)


def p(text="", bold_prefix=None, italic=False, align=None, keep=False):
    para = doc.add_paragraph()
    if align is not None:
        para.alignment = align
    if keep:
        para.paragraph_format.keep_with_next = True
    if bold_prefix and text.startswith(bold_prefix):
        r1 = para.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = para.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        r = para.add_run(text)
        set_run_font(r, italic=italic)
    return para


def h1(text):
    doc.add_heading(text, level=1)


def h2(text):
    doc.add_heading(text, level=2)


def h3(text):
    doc.add_heading(text, level=3)


def bullet(text, level=0):
    para = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    run = para.add_run(text)
    set_run_font(run)
    para.paragraph_format.space_after = Pt(2.5)
    return para


def number(text):
    para = doc.add_paragraph(style="List Number")
    run = para.add_run(text)
    set_run_font(run)
    para.paragraph_format.space_after = Pt(2.5)
    return para


def key_box(title, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE)
    set_cell_margins(cell, 120, 150, 120, 150)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    para = cell.paragraphs[0]
    r = para.add_run(title)
    set_run_font(r, 10.5, bold=True, color=NAVY)
    for line in lines:
        q = cell.add_paragraph(style=None)
        q.paragraph_format.left_indent = Inches(0.12)
        q.paragraph_format.space_after = Pt(2)
        rr = q.add_run("• " + line)
        set_run_font(rr, 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def qa(question, answer, evidence=None):
    # Keep each question and answer in a single non-splitting row so a question is
    # never stranded at the bottom of one page while its answer starts on the next.
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, 100, 120, 105, 120)
    qr = cell.paragraphs[0].add_run("প্রশ্ন: " + question)
    set_run_font(qr, 10, bold=True, color=NAVY)
    ap = cell.add_paragraph()
    ap.paragraph_format.space_after = Pt(1)
    ar = ap.add_run("উত্তর: " + answer)
    set_run_font(ar, 9.7)
    if evidence:
        ep = cell.add_paragraph()
        ep.paragraph_format.space_after = Pt(0)
        er = ep.add_run("Evidence: " + evidence)
        set_run_font(er, 9, italic=True, color=GRAY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def two_col_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        rr = cell.paragraphs[0].add_run(header)
        set_run_font(rr, 9.3, bold=True, color=WHITE)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, text in enumerate(row):
            if i % 2 == 1:
                set_cell_shading(cells[j], PALE)
            set_cell_margins(cells[j])
            rr = cells[j].paragraphs[0].add_run(str(text))
            set_run_font(rr, 9)
    if widths:
        for row in table.rows:
            for j, width in enumerate(widths):
                row.cells[j].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def section_lesson(title, core, explain, focus=None, terms=None, qas=None):
    h2(title)
    h3("মূল বক্তব্য")
    p(core)
    h3("সহজভাবে বুঝি")
    p(explain)
    if terms:
        h3("যে ধারণাগুলো পরিষ্কার রাখতে হবে")
        for term in terms:
            bullet(term)
    if focus:
        key_box("Defense Board Focus", focus)
    if qas:
        h3("সম্ভাব্য Defense প্রশ্ন ও Model Answer")
        for item in qas:
            qa(*item)


def chapter_break():
    doc.add_page_break()


# Title page
for _ in range(3):
    doc.add_paragraph()
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("CAABC-FL Thesis Defense Preparation Guide")
set_run_font(r, 25, bold=True, color=NAVY)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Abstract থেকে Conclusion পর্যন্ত Section Wise বাংলা Lesson এবং Defense Q and A")
set_run_font(sr, 14, bold=True, color=BLUE)

doc.add_paragraph()
project = doc.add_paragraph()
project.alignment = WD_ALIGN_PARAGRAPH.CENTER
pr = project.add_run("Thesis Title")
set_run_font(pr, 11, bold=True, color=GRAY)
project2 = doc.add_paragraph()
project2.alignment = WD_ALIGN_PARAGRAPH.CENTER
pr2 = project2.add_run("CAABC-FL A Hybrid Privacy-Preserving Framework for Secure Multi-Institutional Collaboration")
set_run_font(pr2, 15, bold=True, color=NAVY)

doc.add_paragraph()
purpose = doc.add_paragraph()
purpose.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = purpose.add_run("Final Defense Preparation Edition")
set_run_font(rr, 11, italic=True, color=GRAY)

doc.add_paragraph()
datep = doc.add_paragraph()
datep.alignment = WD_ALIGN_PARAGRAPH.CENTER
dr = datep.add_run("25 September 2026")
set_run_font(dr, 11, bold=True, color=BLUE)

chapter_break()

h1("এই Guide কীভাবে ব্যবহার করবে")
p("এই নথি thesis report-এর বিকল্প নয়; এটি defense preparation-এর জন্য তৈরি একটি structured study guide। প্রতিটি section-এ প্রথমে thesis-এর বক্তব্য, তারপর সহজ ব্যাখ্যা, board কোন জায়গায় প্রশ্ন করতে পারে, এবং concise model answer দেওয়া হয়েছে। প্রথম পাঠে শুধু ‘মূল বক্তব্য’ ও ‘সহজভাবে বুঝি’ অংশ পড়বে। দ্বিতীয় পাঠে equations, assumptions এবং evidence মনে রাখবে। শেষ দুই দিনে প্রতিটি chapter-এর rapid revision এবং final claim boundary অনুশীলন করবে।")
key_box("উত্তর দেওয়ার চার ধাপ", [
    "প্রথম বাক্যে সরাসরি উত্তর দাও।",
    "তারপর ব্যবহৃত method বা assumption এক বাক্যে বলো।",
    "একটি নির্দিষ্ট result, number, experiment বা design evidence দাও।",
    "শেষে limitation বা claim boundary বলো, বিশেষ করে privacy, poisoning এবং deployment প্রশ্নে।",
])
h2("Thesis-এর এক বাক্যের সারাংশ")
p("এই গবেষণা দেখায় যে raw patient record হাসপাতালেই রেখে, patient-level differential privacy এবং CKKS encrypted aggregation ব্যবহার করে, CAABC-FL-এর মাধ্যমে real ciphertext traffic প্রায় 56 শতাংশ কমানো যায় এবং dense encrypted baseline-এর তুলনায় AUROC pre-specified margin-এর মধ্যে ধরে রাখা যায়; তবে AUPRC, poisoning resistance, threshold key custody এবং real multi-host deployment-এর সীমাবদ্ধতাগুলো স্পষ্টভাবে স্বীকার করতে হয়।")
h2("সবচেয়ে গুরুত্বপূর্ণ Claim Boundary")
two_col_table(
    ["যা বলা নিরাপদ", "যা বলা যাবে না"],
    [
        ("CAABC-FL evaluated protocol-এ AUROC-preserving communication trade-off দিয়েছে।", "CAABC-FL সব metric-এ universal SOTA বা clinically superior।"),
        ("Patient-level DP formal accountant অনুযায়ী ε=1.935, δ=10^-5।", "Membership attack AUC 0.5 হওয়াই complete privacy proof।"),
        ("Registered attacks-এ DP membership inference ও reconstruction কঠিন করেছে।", "System সব privacy attack প্রতিরোধ করে।"),
        ("Validation, rollback এবং reputation longitudinal attack সীমিত করেছে।", "Framework poisoning-proof বা formally Byzantine-robust।"),
        ("Separate OpenFHE run native 3-of-3 multiparty decryption দেখিয়েছে।", "Main service fully decentralized, dropout-tolerant threshold CKKS।"),
        ("Single-host Docker service research-system behaviour মাপা হয়েছে।", "Physical multi-hospital production deployment সম্পন্ন হয়েছে।"),
    ],
    [3.25, 3.25],
)

chapter_break()
h1("বিষয়সূচি")
for line in [
    "Lesson 0: Abstract এবং পুরো গবেষণার মানচিত্র",
    "Chapter 1: Introduction",
    "Chapter 2: Literature Review",
    "Chapter 3: Proposed Methodology",
    "Chapter 4: Results and Discussion",
    "Chapter 5: Conclusion",
    "Final Rapid Revision: key numbers, contribution, limitation এবং viva strategy",
]:
    bullet(line)


# ---------------------------------------------------------------------------
# LESSON 0 ABSTRACT
# ---------------------------------------------------------------------------
chapter_break()
h1("Lesson 0 Abstract এবং পুরো গবেষণার মানচিত্র")
p("Abstract হলো পুরো thesis-এর compressed version। Defense board সাধারণত Abstract থেকেই প্রথম প্রশ্ন তৈরি করে, কারণ এখানে problem, method, result এবং conclusion একই জায়গায় পাওয়া যায়। তাই Abstract মুখস্থ না করে এর logical chain বোঝা জরুরি।")

h2("0.1 গবেষণার প্রেক্ষাপট")
p("বিভিন্ন হাসপাতালের কাছে আলাদা রোগীর data থাকে। এক প্রতিষ্ঠানের data দিয়ে model train করলে population variation কম দেখা যায়, কিন্তু raw data কেন্দ্রীভূত করলে privacy, governance, consent এবং institutional ownership-এর সমস্যা তৈরি হয়। Federated learning raw record স্থানান্তর না করে local training করতে দেয়। তবুও federated learning নিজে সম্পূর্ণ privacy solution নয়, কারণ gradient বা model update থেকে patient information অনুমান করা সম্ভব এবং malicious hospital harmful update পাঠাতে পারে।")

h2("0.2 গবেষণার মূল সমস্যা")
p("এই thesis চারটি coupled problem একসঙ্গে সমাধান করার চেষ্টা করেছে: released model থেকে patient privacy রক্ষা, coordinator-এর কাছ থেকে individual update গোপন রাখা, encrypted update-এর communication cost কমানো, এবং invalid বা malicious action-এর জন্য accountability তৈরি করা। গুরুত্বপূর্ণ insight হলো, একটি control অন্য control-কে বাধাগ্রস্ত করতে পারে। উদাহরণস্বরূপ, encryption update গোপন রাখলেও অনেক robust aggregator individual update geometry না দেখতে পারলে কাজ করতে পারে না।")

h2("0.3 প্রস্তাবিত সমাধান")
p("Framework-টি local patient-level DP training, CKKS encrypted aggregation, CAABC-FL communication reduction, signed and authenticated submission, independent aggregate recomputation, validation and rollback, reputation control, model registry এবং permissioned audit witness একত্র করে। CAABC-FL model update-কে 256-coordinate block-এ ভাগ করে, shared ciphertext-aligned support নির্বাচন করে এবং omitted information error-feedback residual-এ রাখে। ফলে client-specific sparsity প্রকাশ না করে একই CKKS slot layout-এ encrypted summation করা যায়।")

h2("0.4 Evaluation এবং principal result")
p("Evaluation PhysioNet 2019-এর 40,336 patient এবং দুইটি natural hospital partition-এর উপর করা হয়েছে। Core predictive comparisons দশ paired seed এবং computationally expensive SOTA, poisoning, membership inference, reconstruction ও ablation experiments পাঁচ paired seed ব্যবহার করেছে। Frozen private configuration-এ ε=1.935 এবং δ=10^-5। CAABC-FL dense DP-FedAvg-CKKS-এর 3.314 MB প্রতি round encrypted upload কমিয়ে 1.458 MB করেছে, অর্থাৎ 55.998 শতাংশ measured ciphertext reduction। AUROC difference +0.00017 এবং 95 শতাংশ confidence interval [-0.00212, +0.00247], যা -0.015 non-inferiority margin-এর উপরে ছিল।")
p("তবে AUPRC 0.00926 কমেছে এবং calibration ও কিছু threshold-dependent metric-ও খারাপ হয়েছে। তাই সঠিক conclusion হলো CAABC-FL একটি AUROC-preserving encrypted communication trade-off দিয়েছে; এটি সব metric-এ superior নয়। DP membership-inference attack-কে প্রায় chance-level করেছে এবং gradient reconstruction RMSE বাড়িয়েছে, কিন্তু weighted encrypted mean শক্তিশালী model-replacement attack-এর বিরুদ্ধে দুর্বল ছিল।")

key_box("Abstract থেকে মনে রাখবে", [
    "Dataset: PhysioNet 2019, 40,336 patients, two natural hospital partitions.",
    "Privacy: patient-level DP, ε=1.935 এবং δ=10^-5.",
    "Method: CAABC-FL with shared aligned blocks, adaptive refresh এবং error feedback.",
    "Communication: 3.314 MB থেকে 1.458 MB, reduction 55.998%.",
    "Utility: AUROC non-inferiority passed; AUPRC loss আছে.",
    "Security boundary: attack-specific evidence, no universal privacy or poisoning-proof claim.",
])

h2("0.5 Abstract ভিত্তিক Defense প্রশ্ন")
qa("আপনাদের গবেষণার novelty এক বাক্যে কী?", "CAABC-FL encrypted federated learning-এ adaptive compression-কে CKKS packing boundary-এর সঙ্গে align করে, যাতে logical sparsity বাস্তব serialized ciphertext saving-এ রূপ নেয়; একই system-এ patient-level privacy, authenticated aggregation এবং accountable model release-ও মূল্যায়ন করা হয়েছে।", "55.998% encrypted-wire reduction এবং AUROC non-inferiority pass।")
qa("Federated learning থাকলেই privacy কেন নিশ্চিত হয় না?", "কারণ raw data local থাকলেও gradients, model deltas বা final model থেকে membership বা feature information ফাঁস হতে পারে; server update manipulate করতে পারে এবং malicious participant poisoning করতে পারে। তাই DP, encryption, authentication এবং governance আলাদা প্রয়োজন।")
qa("আপনার main contribution algorithm নাকি framework?", "দুই স্তরের contribution আছে। Algorithmic contribution হলো CAABC-FL; systems contribution হলো privacy, encrypted aggregation, verification এবং governance-কে একটি executable research framework-এ যুক্ত করা। তবে strongest measured contribution হলো real encrypted communication reduction with bounded AUROC loss।")
qa("আপনারা কি state of the art?", "আমরা universal SOTA claim করি না। Same-protocol পাঁচ-seed comparison-এ CAABC-FL relevant baselines-এর বিরুদ্ধে AUROC ও AUPRC non-inferiority margins পাস করেছে এবং একমাত্র substantial measured encrypted-wire reduction দেখিয়েছে। Different dataset বা threat model-এর paper-এর সঙ্গে headline number দিয়ে superiority দাবি করা হয়নি।")
qa("সবচেয়ে দুর্বল result কোনটি?", "AUPRC dense encrypted baseline-এর তুলনায় 0.00926 কমেছে এবং strong one-round model-replacement attack weighted encrypted mean-কে ক্ষতিগ্রস্ত করেছে। এই negative results আমাদের claim সীমিত করে এবং future encrypted robust aggregation-এর প্রয়োজন দেখায়।")


# ---------------------------------------------------------------------------
# CHAPTER 1
# ---------------------------------------------------------------------------
chapter_break()
h1("Chapter 1 Introduction")
p("Chapter 1 পাঠককে problem domain থেকে research objective পর্যন্ত নিয়ে যায়। Board এখানে দেখবে তোমরা সমস্যাটি বাস্তবভাবে বুঝেছ কি না, objectives measurable কি না এবং claimed contribution research problem-এর সঙ্গে align করে কি না।")

section_lesson(
    "1.0 Introduction",
    "Multi-institutional healthcare learning-এর সম্ভাবনা বড়, কারণ বিভিন্ন হাসপাতালের combined experience থেকে বেশি representative model তৈরি হতে পারে। কিন্তু patient record এক জায়গায় জমা করা privacy, regulation, ownership এবং breach risk বাড়ায়। এই thesis secure collaborative learning-এর জন্য privacy, confidentiality, communication এবং trust একসঙ্গে বিবেচনা করে।",
    "ধরো দুই হাসপাতালের রোগীর ধরন আলাদা। আলাদা model করলে প্রত্যেকটি model নিজের population-এ আটকে যেতে পারে। Data সরাসরি share করলে risk বাড়ে। Federated learning মাঝামাঝি পথ দেয়: data local, model update shared। কিন্তু update নিজেও sensitive এবং expensive; এখানেই thesis-এর প্রয়োজন।",
    ["FL-কে complete privacy solution নয়, collaboration mechanism হিসেবে সংজ্ঞায়িত করবে।", "Problemটি healthcare হলেও contribution secure cross-silo collaboration-এর বৃহত্তর context-এও প্রযোজ্য।"],
    ["Cross-silo FL: অল্পসংখ্যক পরিচিত institution দীর্ঘমেয়াদে সহযোগিতা করে।", "Cross-device FL: বিপুল সংখ্যক অনিশ্চিত mobile বা edge device অংশ নেয়।"],
    [
        ("এই কাজ cross-silo কেন?", "কারণ participant হলো পরিচিত hospital বা institution, তাদের identity, policy, record volume এবং long-lived relationship আছে। Mobile-device scale বা random participation এই study-এর target নয়।"),
        ("Healthcare context কেন কঠিন?", "Data highly sensitive, class imbalance ও non-IID distribution সাধারণ, clinical false negative costly, এবং governance ও audit requirement সাধারণ application-এর চেয়ে কঠোর।"),
    ],
)

section_lesson(
    "1.1 Background of the Study",
    "Distributed machine learning, federated averaging, differential privacy, homomorphic encryption এবং secure aggregation ধাপে ধাপে privacy-preserving collaboration-এর foundation তৈরি করেছে। Recent work communication compression, Byzantine defence এবং accountable governance যোগ করেছে, কিন্তু সব component একই assumption-এর অধীনে একত্র করা কঠিন।",
    "Field-এর evolution তিন ধাপে বোঝা যায়: প্রথমে data centralize করে ML; তারপর data local রেখে FL; এরপর FL update-এর privacy, integrity এবং cost সমাধানে DP, HE, compression ও governance। Thesis এই তৃতীয় ধাপের integration problem নিয়ে কাজ করে।",
    ["Historical milestones বলতে invention-এর তালিকা নয়; problem-to-solution evolution দেখাবে।", "HE confidentiality দেয়, DP output leakage bound করে; দুটির ভূমিকা এক নয়।"],
    ["FedAvg 2017 সালে FL optimization-এর canonical baseline।", "DP patient contribution সীমিত করে; CKKS approximate real-number computation সম্ভব করে।", "Modern challenge হলো privacy, robustness এবং efficiency simultaneously অর্জন করা।"],
    [
        ("DP এবং HE দুটোই কেন দরকার?", "HE transport ও aggregation-এর সময় update গোপন রাখে, কিন্তু decrypted aggregate বা final model-এর information leakage সীমিত করে না। DP একজন patient-এর influence bound করে। তাই তারা complementary।"),
        ("Blockchain কি অপরিহার্য?", "না। আমাদের requirement হলো tamper-evident, independently witnessed audit। Implemented witness plane permissioned signed checkpoint দেয়; এটিকে public blockchain বা Byzantine consensus বলা হয়নি।"),
    ],
)

section_lesson(
    "1.1.1 State of the Art in the Modern Context",
    "বর্তমান research private aggregation, record-level DP, homomorphic computation, sparse communication এবং Byzantine defence-এ শক্তিশালী individual solutions দিয়েছে। কিন্তু encrypted aggregation individual update লুকিয়ে দেয়, আবার Krum, trimmed mean বা FLTrust-এর মতো defence update geometry বা trusted root data চায়। একইভাবে Top-k logical sparsity CKKS-এ dense re-encoding হলে real network saving দেয় না।",
    "State of the art বিচার করতে শুধু accuracy দেখা যাবে না। একই সঙ্গে জানতে হবে privacy unit কী, server কী দেখে, attack model কী, transmitted bytes কীভাবে মাপা হয়েছে এবং defence-এর trust assumption কী। Thesis এই comparison dimensions ব্যবহার করেছে।",
    ["Related systems-এর result direct ranking নয় যদি dataset, privacy budget বা threat model আলাদা হয়।", "Same-protocol reproduction strongest comparative evidence; literature comparison contextual evidence।"],
    ["Privacy-robustness compatibility gap", "Logical compression বনাম encrypted-wire reduction", "Formal guarantee বনাম empirical attack resistance"],
    [
        ("Modern literature-এর প্রধান gap কী?", "একই evaluated protocol-এ record-level privacy, real encrypted aggregation, ciphertext-aware communication reduction, poisoning evaluation এবং accountable release governance-এর সমন্বয় খুব সীমিত।"),
        ("Plaintext robust aggregator সরাসরি ব্যবহার করেননি কেন?", "কারণ coordinator যদি individual update vector না দেখে, coordinate median, Krum distance বা trusted-root cosine score সরাসরি হিসাব করতে পারে না; তা করলে confidentiality assumption বদলে যায়।"),
    ],
)

section_lesson(
    "1.1.2 Importance and Relevance of the Topic",
    "Secure institutional collaboration ছোট hospital-কে broader population knowledge ব্যবহার করতে সাহায্য করতে পারে এবং raw data transfer কমাতে পারে। এর value healthcare ছাড়াও finance, research consortium, smart infrastructure এবং regulated analytics-এ আছে। তবে benefit তখনই অর্থবহ যখন privacy claim, communication overhead এবং accountability বাস্তব measurement দিয়ে সমর্থিত হয়।",
    "একটি hospital একা rare case কম দেখে। Collaboration model-এর generalization বাড়াতে পারে। কিন্তু patient trust হারালে প্রযুক্তিগত লাভ মূল্যহীন। তাই study-এর relevance শুধু ভালো classifier নয়, responsible collaboration infrastructure।",
    ["Social benefit এবং technical feasibility দুটোই বলবে।", "Research system-কে clinical decision maker হিসেবে উপস্থাপন করবে না।"],
    ["Patient privacy", "Institutional autonomy", "Cross-site generalization", "Auditability and accountability"],
    [
        ("এই research সমাজে কী impact ফেলতে পারে?", "Secure collaboration future clinical decision-support model তৈরি করতে সাহায্য করতে পারে, বিশেষ করে যেখানে এক প্রতিষ্ঠানের data সীমিত। একই সঙ্গে raw record custody বজায় রেখে institutional accountability বাড়ায়।"),
    ],
)

section_lesson(
    "1.2 Research Problem and Its Context",
    "বর্তমান FL system সাধারণত privacy, communication বা robustness-এর একাংশ সমাধান করে। Raw data local থাকলেও updates leak করতে পারে; CKKS ciphertext বড়; client-specific sparsity aggregation ও privacy-র সঙ্গে conflict করে; malicious update encrypted mean-কে corrupt করতে পারে; এবং decryption ও model release-এর accountability দুর্বল হতে পারে।",
    "মূল problem শুধু ‘কীভাবে model train করব’ নয়। Problem হলো ‘কীভাবে বহু প্রতিষ্ঠানকে এমনভাবে train করাব যাতে raw data না যায়, patient influence সীমিত থাকে, server individual update না দেখে, encrypted upload কমে এবং release কে অনুমোদন করেছে তা যাচাই করা যায়’।",
    ["Problem statement-এ solution-এর নামের বদলে unmet requirements দেখাবে।", "চারটি dimension: privacy, confidentiality, efficiency, integrity/governance।"],
    ["Non-IID institutional data", "Update leakage", "Ciphertext expansion", "Poisoning and auditability"],
    [
        ("আপনার research problem কি খুব বড় নয়?", "এটি modular engineering problem হিসেবে bounded করা হয়েছে। আমরা নির্দিষ্ট dataset, model, DP configuration, CKKS parameters, attack families এবং research deployment-এর মধ্যে integrated feasibility ও trade-off evaluate করেছি। Universal clinical deployment দাবি করিনি।"),
        ("Problemটির measurable form কী?", "Primary evidence ছিল dense encrypted reference-এর তুলনায় AUROC non-inferiority এবং 40 শতাংশের বেশি actual encrypted-wire reduction; privacy, poisoning ও governance secondary but essential evaluation families।"),
    ],
)

section_lesson(
    "1.3 Motivation",
    "Research-এর motivation আসে patient confidentiality, institutional data sovereignty, encrypted communication cost এবং malicious collaboration risk থেকে। Technical motivation হলো protective mechanisms আলাদাভাবে সফল হলেও combined system-এ তাদের assumptions conflict করে।",
    "Hospital data valuable কিন্তু share করা sensitive। FL attractive, তবে একটি বড় encrypted update বহু round পাঠালে cost বেড়ে যায়। আবার compression ভুলভাবে করলে CKKS ciphertext size কমে না। এই বাস্তব engineering mismatch CAABC-FL-এর motivation।",
    ["Motivation-এ personal interest বললে health-data privacy ও trustworthy AI-এর প্রতি আগ্রহের কথা বলতে পারো।", "‘Accuracy বাড়ানো’ একমাত্র motivation নয়।"],
    ["Privacy-preserving collaboration", "Measured network efficiency", "Adversarial resilience", "Responsible governance"],
    [
        ("CAABC-FL তৈরির direct motivation কী ছিল?", "Existing sparse methods retained coordinate কমালেও dense CKKS transport-এ ciphertext count অপরিবর্তিত থাকতে পারে। আমরা compression decision-কে ciphertext block layout-এর সঙ্গে align করতে চেয়েছি।"),
    ],
)

section_lesson(
    "1.4 Applications and Practical Use Cases",
    "Potential use cases-এর মধ্যে inter-hospital risk modelling, rare-disease research, public-health surveillance, multi-centre quality improvement এবং regulated research consortium রয়েছে। Framework-এর governance layer model release এবং audit review-এ সাহায্য করতে পারে।",
    "এটি bedside diagnostic device নয়। বরং research-stage collaborative training platform। ভবিষ্যতে proper external validation, clinical governance এবং regulatory approval থাকলে trained model decision support-এ ব্যবহৃত হতে পারে।",
    ["Current capability এবং future possibility আলাদা করবে।", "Inference path-এর consent, purpose binding এবং minimal disclosure governance use case-এর অংশ।"],
    ["Cross-hospital sepsis research", "Federated cohort analytics", "Auditable consortium model training"],
    [
        ("এখনই hospital-এ deploy করা যাবে?", "না। System research prototype-এর চেয়ে উন্নত executable research system, কিন্তু clinical deployment-এর জন্য external prospective validation, usability study, regulatory review, security audit এবং physical multi-host testing বাকি।"),
    ],
)

section_lesson(
    "1.5 Research Objectives",
    "General objective হলো secure, privacy-preserving এবং communication-efficient multi-institutional collaborative learning framework develop ও evaluate করা। Specific objectives তিনটি: integrated privacy and trust architecture তৈরি; protected model exchange-এর communication overhead কমানো; এবং benign ও adversarial condition-এ relevant baselines-এর সঙ্গে comprehensive evaluation করা।",
    "Objective-গুলো generalized রাখা হয়েছে যাতে technique নাম objective-কে অযথা narrow না করে। Methodology chapter পরে বলে কীভাবে objective achieve করা হয়েছে। Results chapter objective-wise evidence দেয়।",
    ["Objective 1 = system integration", "Objective 2 = communication method", "Objective 3 = empirical evaluation"],
    ["Objectives measurable এবং contribution-এর সঙ্গে traceable।", "Metric-এর নাম objective-এ না থাকলেও evaluation plan-এ স্পষ্ট।"],
    [
        ("তিনটি objective অর্জিত হয়েছে কি?", "Research-system scope-এ হয়েছে। Integrated architecture implemented; CAABC-FL real encrypted traffic কমিয়েছে; utility, privacy, poisoning, systems এবং governance evaluation করা হয়েছে। তবে production threshold deployment ও external clinical validation objective-এর বাইরে future work।"),
    ],
)

section_lesson(
    "1.6 Research Methodology Overview",
    "Study design-science এবং controlled quantitative experiment একত্র করেছে। প্রথমে literature ও threat analysis থেকে requirements নেওয়া হয়েছে; পরে architecture ও CAABC-FL artifact তৈরি; তারপর frozen protocol-এ healthcare data, paired seeds, baselines, privacy attacks, poisoning attacks এবং system tests চালানো হয়েছে; শেষে statistical confidence interval ও claim boundary দিয়ে interpretation করা হয়েছে।",
    "Design science মানে শুধু theory নয়, একটি working artifact তৈরি এবং evidence দিয়ে evaluate করা। Paired design একই seed, split ও configuration-এ methods তুলনা করে variance কমায়।",
    ["Build, evaluate, interpret, bound claims", "Training/validation/test leakage control", "Serialized ciphertext measurement"],
    ["Retrospective experimental study", "Natural institutional partition", "Paired multi-seed comparison"],
    [
        ("কেন design-science approach?", "কারণ research question একটি integrated engineering artifact নির্মাণ এবং তার utility, privacy, security ও system behaviour মূল্যায়ন চায়। শুধু survey বা predictive modelling যথেষ্ট ছিল না।"),
        ("কেন paired seeds?", "একই random seed ও data split-এ দুই method তুলনা করলে between-seed noise বাদ দিয়ে method effect আরও নির্ভুলভাবে estimate করা যায়।"),
    ],
)

section_lesson(
    "1.7 Research Contribution",
    "Contribution পাঁচ ভাগে ব্যাখ্যা করা যায়: conceptual integration of privacy-confidentiality-efficiency-governance; CAABC-FL algorithm; executable multi-service research system; rigorous evaluation protocol; এবং measured empirical evidence। Strongest empirical evidence হলো 55.998 শতাংশ real encrypted-wire reduction with AUROC non-inferiority।",
    "Contribution মানে feature list নয়। যা নতুন, কেন দরকার এবং কী evidence আছে - এই তিনটি বলতে হবে। CAABC-FL-এর novelty block compression নয় একা; novelty হলো adaptive shared support-কে CKKS packing ও measured wire cost-এর সঙ্গে যুক্ত করা।",
    ["Contribution এবং result আলাদা কিন্তু connected।", "Negative result-ও contribution: encryption poisoning solve করে না এবং plaintext defence confidentiality assumption ভাঙে।"],
    ["Theoretical framing", "Algorithmic design", "System artifact", "Methodological rigor", "Empirical evidence"],
    [
        ("সবচেয়ে original contribution কোনটি?", "CAABC-FL-এর ciphertext-aligned adaptive common-block selection with error feedback, কারণ এটি sparse learning decision-কে encrypted transport representation-এর সঙ্গে co-design করে।"),
        ("Framework contribution-এর evidence কী?", "Hospital, coordinator, authority এবং witness services; signed envelopes; mTLS; independent aggregate recomputation; persistent registry; rollback এবং witnessed audit failure cases successfully exercised হয়েছে।"),
    ],
)

section_lesson(
    "1.8 Financial Planning and Timeline",
    "Study open-source software, public research data এবং personal workstation ব্যবহার করেছে; তাই major procurement cost ছিল না। Timeline literature review, architecture, implementation, experiment, analysis এবং writing phase-এ ভাগ করা হয়েছে।",
    "Low budget মানে zero resource cost নয়। Compute time, electricity, storage এবং researcher time আছে, কিন্তু direct project expenditure সীমিত। Gantt chart কাজের dependency ও completion tracking দেখায়।",
    ["Open-source does not mean cost-free; direct financial cost low।", "Timeline iterative ছিল: result দেখে protocol refine, তারপর confirmatory run।"],
    ["Literature and design", "Implementation", "Evaluation", "Writing and revision"],
    [
        ("Budget কম হওয়ায় experiment দুর্বল হয়েছে?", "Public dataset ও open-source stack reproducibility বাড়িয়েছে, তবে single-workstation constraint physical multi-host এবং hardware-attestation experiment সীমিত করেছে। এই limitation report-এ stated।"),
    ],
)

section_lesson(
    "1.9 Organization of the Dissertation",
    "Chapter 1 problem ও objective দেয়; Chapter 2 theory, related studies এবং gap বিশ্লেষণ করে; Chapter 3 system ও experiment design ব্যাখ্যা করে; Chapter 4 measured result এবং interpretation দেয়; Chapter 5 contribution, limitation ও future work সংক্ষেপ করে।",
    "Organization section পাঠককে roadmap দেয়। Defense-এও একই roadmap অনুসরণ করলে উত্তর coherent হয়: why, what exists, what we built, what we found, what remains।",
    ["Why -> Gap -> Method -> Evidence -> Boundary"],
    ["প্রতিটি chapter পরের chapter-এর প্রয়োজন তৈরি করে।"],
    [
        ("Literature review-এর gap কীভাবে methodology-তে গেছে?", "Privacy-robustness incompatibility threat model ও post-aggregation governance তৈরি করেছে; logical-to-ciphertext gap CAABC common block design তৈরি করেছে; trust gap independent recomputation ও witness audit তৈরি করেছে।"),
    ],
)

h2("Chapter 1 Rapid Revision")
key_box("60 সেকেন্ডে Chapter 1", [
    "Problem: raw data local থাকলেও update leakage, ciphertext cost, poisoning এবং trust সমস্যা থাকে।",
    "Aim: secure, private, efficient এবং accountable multi-institutional collaboration।",
    "Three objectives: integrated architecture, communication reduction, comprehensive evaluation।",
    "Novelty: CAABC-FL plus integrated executable system and rigorous evidence।",
    "Boundary: research system, not clinical product or universal SOTA।",
])


# ---------------------------------------------------------------------------
# CHAPTER 2
# ---------------------------------------------------------------------------
chapter_break()
h1("Chapter 2 Literature Review")
p("Literature review-এর কাজ paper-এর তালিকা দেওয়া নয়; existing knowledge synthesize করে design decision justify করা। এই chapter-এ optimization, privacy, encryption, communication, poisoning defence এবং governance আলাদা stream হিসেবে দেখা হয়েছে, তারপর তাদের compatibility gap বিশ্লেষণ করা হয়েছে।")

section_lesson(
    "2.1 Introduction",
    "Review recent peer-reviewed journals, conferences, standards-oriented technical literature এবং foundational papers cover করে। Search themes ছিল healthcare federated learning, differential privacy, homomorphic encryption, secure aggregation, communication compression, poisoning defence, audit এবং threshold governance। Inclusion-এর মূল শর্ত ছিল method, assumptions বা empirical evidence research problem-এর সঙ্গে relevant হওয়া।",
    "Review-এর scope ইচ্ছাকৃতভাবে broad, কারণ thesis একটি hybrid framework। শুধু FL paper পড়লে encryption cost বোঝা যাবে না; শুধু cryptography paper পড়লে clinical metric ও poisoning behaviour বোঝা যাবে না।",
    ["Search strategy transparent হওয়া উচিত, কিন্তু PRISMA systematic review দাবি করা যাবে না যদি full screening protocol না থাকে।", "Foundational পুরোনো paper এবং recent integrated work উভয় দরকার।"],
    ["Theme-based synthesis", "Assumption-aware comparison", "Direct vs contextual evidence"],
    [
        ("এটি systematic literature review কি?", "এটি structured thematic review। Databases, keywords এবং inclusion logic documented, কিন্তু formal PRISMA meta-analysis বা exhaustive systematic review হিসেবে claim করা হয়নি।"),
        ("কেন 2017-এর আগের paper আছে?", "Recent ten-year scope-এর মধ্যে modern FL literature প্রধান, তবে differential privacy ও cryptography-এর foundational definitions বোঝাতে earlier canonical sources প্রয়োজন।"),
    ],
)

section_lesson(
    "2.2.1 Cross-Silo Federated Learning",
    "Cross-silo FL-এ কয়েকটি known institution একই global model train করে। প্রতিটি site local objective optimize করে এবং coordinator weighted aggregation করে। Data distribution non-IID হতে পারে এবং institution-এর sample size ভিন্ন হয়।",
    "এই thesis-এ দুই natural hospital system artificial random client-এ ভাঙা হয়নি। তাই institutional heterogeneity বাস্তব থাকে। Client identity known হওয়ায় authentication, reputation এবং consortium policy ব্যবহার করা সম্ভব।",
    ["Global objective: F(theta)=sum alpha_k F_k(theta)", "alpha_k সাধারণত n_k divided by total n", "Raw data কখনও coordinator-এ যায় না"],
    ["Cross-silo assumptions cross-device-এর মতো নয়।", "Natural partition external validity বাড়ায়, কিন্তু দুই institution generalization সীমিত করে।"],
    [
        ("Natural partition কেন গুরুত্বপূর্ণ?", "Random partition করলে site-specific prevalence, workflow ও measurement difference মুছে যেতে পারে। Natural partition বাস্তব cross-institution heterogeneity retain করে।"),
    ],
)

section_lesson(
    "2.2.2 Federated Optimization under Heterogeneity",
    "FedAvg local SGD updates weighted-average করে। FedProx proximal term দিয়ে client drift কমায়; SCAFFOLD control variate ব্যবহার করে; FedNova local-step mismatch normalize করে। এগুলো optimization problem সমাধান করে, privacy বা encrypted transport স্বয়ংক্রিয়ভাবে দেয় না।",
    "Non-IID data-তে hospital A-এর gradient এবং hospital B-এর gradient ভিন্ন দিকে যেতে পারে। Local training বেশি করলে communication কমে কিন্তু client drift বাড়তে পারে। তাই optimizer comparison utility ceiling ও heterogeneity effect বোঝায়।",
    ["Optimization baseline এবং privacy baseline আলাদা।", "FedAvg high utility দিলেও ε=0 নয়; privacy mechanism না থাকলে formal DP নেই।"],
    ["Client drift", "Local steps vs communication", "Weighted aggregation"],
    [
        ("FedProx বা SCAFFOLD-কে proposed method বলা হয়নি কেন?", "তারা established optimization baselines। আমাদের novelty optimizer replacement নয়; protected communication এবং integrated governance।"),
        ("FedAvg বেশি AUROC পেয়েছে, তাহলে CAABC কেন দরকার?", "FedAvg privacy ও update confidentiality দেয় না। Fair comparison হলো dense DP-FedAvg-CKKS-এর সঙ্গে, কারণ সেটি একই protection path রাখে।"),
    ],
)

section_lesson(
    "2.2.3 Differential Privacy",
    "Differential privacy neighbouring datasets-এর output distribution কাছাকাছি রাখে। Patient-level adjacency-তে একটি patient-এর সম্পূর্ণ contribution যোগ বা বাদ হলে mechanism-এর behaviour bounded থাকে। DP-SGD per-sample gradient clip করে এবং Gaussian noise যোগ করে; accountant repeated steps-এর privacy loss compose করে।",
    "Clipping এক রোগীর influence C-এর মধ্যে রাখে। Noise attackers-এর signal blur করে। Epsilon ছোট হলে সাধারণত stronger privacy, কিন্তু utility cost বেশি হতে পারে। Delta একটি small failure probability term। এই study-তে final recorded budget ε=1.935, δ=10^-5।",
    ["DP protects released information, not transport confidentiality।", "Post-processing of already privatized update নতুন privacy budget consume করে না।", "Attack AUC near 0.5 formal proof নয়।"],
    ["Neighbouring datasets", "Clipping norm C=1.0", "Noise multiplier sigma=1.2", "RDP accounting"],
    [
        ("Patient-level DP আর sample-level DP-এর পার্থক্য কী?", "Patient-level DP-তে একজন patient-এর সব row বা time point এক privacy unit। ICU time-series-এ একই patient-এর বহু measurement থাকায় row-level protection যথেষ্ট নয়।"),
        ("ε=1.935 কি ভালো?", "এটি stated accountant, sampling এবং adjacency-এর অধীনে moderate formal budget। একে absolute good বা bad বলা যাবে না; utility, policy এবং comparable deployments-এর context-এ interpret করতে হয়।"),
        ("Compression কি privacy বাড়ায়?", "শুধু sparsity privacy guarantee নয়। CAABC support privatized update-এর post-processing, কিন্তু formal privacy DP accountant থেকেই আসে।"),
    ],
)

section_lesson(
    "2.2.4 Homomorphic Encryption Secure Aggregation and Threshold Governance",
    "CKKS approximate real-number vectors-এর encrypted addition ও scalar multiplication support করে। Coordinator public context দিয়ে ciphertext aggregate করতে পারে কিন্তু secret key ছাড়া plaintext update দেখতে পারে না। Secure aggregation একই লক্ষ্য ভিন্ন protocol-এ দেয়। Threshold governance secret key বা decryption authority এক actor-এর হাতে না রাখার নীতি।",
    "DP প্রশ্ন করে final output একজন patient সম্পর্কে কতটা বদলায়। CKKS প্রশ্ন করে computation চলাকালে server update পড়তে পারে কি না। Threshold system প্রশ্ন করে কে decrypt করতে পারবে। এগুলো তিনটি আলাদা layer।",
    ["CKKS approximate arithmetic; tiny numerical error expected।", "Polynomial modulus degree, coefficient chain এবং scale security, precision ও cost নির্ধারণ করে।", "Main TenSEAL service centrally keyed; OpenFHE threshold experiment separate।"],
    ["Public-key coordinator", "Packed slots", "Ciphertext expansion", "3-of-3 multiparty demonstration"],
    [
        ("CKKS কেন, BFV নয়?", "Model update real-valued। CKKS approximate real arithmetic এবং packed vector operations-এর জন্য suitable; BFV/BGV exact integer arithmetic-এর জন্য বেশি natural।"),
        ("Threshold decryption implemented হয়েছে?", "Separate OpenFHE test native 3-of-3 multiparty decryption দেখিয়েছে, কিন্তু main network service-এ arbitrary t-of-n বা dropout-tolerant custodian integration হয়নি।"),
        ("CKKS security parameter কী ছিল?", "Polynomial modulus degree 8192, coefficient-modulus chain 60-40-40-60 bits এবং global scale 2^40।"),
    ],
)

section_lesson(
    "2.2.5 Communication-Efficient Federated Learning",
    "Communication compression সাধারণত sparsification, quantization, sketching বা periodic communication ব্যবহার করে। Error feedback transmitted না হওয়া residual পরে পুনরায় যোগ করে optimization bias কমাতে সাহায্য করে। Encryption-এর সঙ্গে efficiency বিচার করতে serialized ciphertext bytes মাপা জরুরি।",
    "Top-k plaintext vector-এ 10 শতাংশ coordinate রাখলেও, privacy রক্ষার জন্য sparse vector dense CKKS ciphertext-এ re-encode করলে network packet প্রায় আগের মতোই থাকে। CAABC shared block support ব্যবহার করে ciphertext chunk-এর সংখ্যাই কমায়।",
    ["Logical reduction != wire reduction", "Common support aggregation-compatible", "Refresh message-ও wire cost-এর মধ্যে গণনা করা হয়েছে"],
    ["Top-k", "Quantization", "Error feedback", "Ciphertext packing alignment"],
    [
        ("CAABC Top-k থেকে কীভাবে আলাদা?", "Top-k client-specific coordinates বেছে নেয়। CAABC contiguous block score থেকে consortium-wide common support নেয়, যা CKKS slot alignment ও encrypted summation বজায় রেখে fewer ciphertext chunks পাঠায়।"),
        ("Error feedback কেন?", "Omitted private update residual-এ রেখে পরের round-এর effective update-এ যোগ করা হয়, যাতে বারবার বাদ পড়া information স্থায়ীভাবে হারিয়ে না যায়।"),
    ],
)

section_lesson(
    "2.2.6 Threat Model Privacy Integrity and Availability",
    "Privacy threat-এর মধ্যে update leakage, membership inference এবং gradient reconstruction; integrity threat-এর মধ্যে poisoning, backdoor, replay, forgery এবং aggregate substitution; availability threat-এর মধ্যে denial of progress, witness outage ও over-sensitive rollback রয়েছে।",
    "Security control threat-specific। Encryption eavesdropping কমায়, signature impersonation আটকায়, nonce replay আটকায়, DP patient influence bound করে, আর validation harmful model release থামায়। একটি control দিয়ে সব threat solve হয় না।",
    ["Adversary capability স্পষ্ট না হলে security claim অর্থহীন।", "Honest-but-curious coordinator protocol follow করে কিন্তু accessible message থেকে infer করতে চেষ্টা করে।"],
    ["Confidentiality", "Integrity", "Availability", "Authentication", "Accountability"],
    [
        ("Threat model-এ কী exclude করা হয়েছে?", "Physical host compromise, library side channel, centralized authority compromise, majority witness collusion এবং formally audited custom RNG exclude করা হয়েছে।"),
        ("Rollback security বাড়ালেও সমস্যা কী?", "Rollback harmful release আটকাতে পারে, কিন্তু attacker repeated trigger করলে learning progress থামাতে পারে। তাই integrity এবং availability joint evaluation দরকার।"),
    ],
)

section_lesson(
    "2.2.7 Robust Aggregation and Trust Management",
    "Coordinate median, trimmed mean, Krum, RFA এবং FLTrust malicious update reduce করতে পারে, কিন্তু তাদের effectiveness attacker fraction ও data assumptions-এর উপর নির্ভর করে। বেশিরভাগ method individual update geometry অথবা trusted clean root dataset চায়। Reputation temporal behaviour থেকে participant weight বা quarantine policy চালায়।",
    "Encryption-এর conflict এখানেই। Sum-only server individual vector দেখতে না পারলে কোন update outlier তা সরাসরি বলা কঠিন। আমাদের framework তাই encrypted mean-এর পরে validation, rollback ও reputation ব্যবহার করেছে, যদিও এটি formal Byzantine aggregation নয়।",
    ["Plaintext defence stronger হতে পারে কিন্তু confidentiality trade-off আছে।", "Reputation Sybil-proof নয় যদি identity governance দুর্বল হয়।"],
    ["Byzantine bound", "Trusted root", "Quarantine recall", "Rollback rate"],
    [
        ("Multi-Krum ব্যবহার করলে poisoning solve হতো না?", "Registered one-round tests-এ Multi-Krum strong ছিল, কিন্তু এটি pairwise distance-এর জন্য individual plaintext updates চায়। Sum-only CKKS privacy path-এর direct drop-in replacement নয়।"),
        ("আপনার reputation system কি robust aggregation?", "না। এটি governance control যা past validation evidence দিয়ে future influence বা participation সীমিত করে; formal robust estimator নয়।"),
    ],
)

section_lesson(
    "2.2.8 Permissioned Ledgers and Independent Verification",
    "Permissioned audit system known actors-এর signed event, digest, timestamp এবং chain continuity সংরক্ষণ করে। Independent verifier coordinator-এর aggregate পুনরায় compute করে omission বা substitution detect করতে পারে। Ledger correctness proof নয়; এটি sequence এবং authorization-এর evidence।",
    "আমাদের witness plane তিনটি separately keyed persistent witness ব্যবহার করে এবং দুইটি receipt quorum চায়। একটি witness বন্ধ থাকলেও progress হয় এবং ফিরে এলে reconciliation হয়। কিন্তু এটি blockchain consensus নয়।",
    ["Audit log tamper-evidence দেয়, truth guarantee নয়।", "Independent recomputation zero-knowledge proof নয়।"],
    ["Digest binding", "Signed envelope", "Two-of-three witness receipts", "Model registry"],
    [
        ("Blockchain ব্যবহার করেছেন?", "না। External permissioned witness ledger ব্যবহার করেছি। এটি signed hash-chained checkpoints দেয়, কিন্তু proof-of-work, token বা BFT blockchain consensus নয়।"),
        ("Coordinator ভুল aggregate দিলে কী হয়?", "Authority original signed envelopes reverify করে deterministic aggregate পুনর্গঠন করে এবং digest compare করে। Mismatch হলে decryption/release reject হয়।"),
    ],
)

section_lesson(
    "2.3 Related Studies",
    "Related work পাঁচ stream-এ organize করা হয়েছে: healthcare FL demonstration; privacy-enhanced learning; encrypted and communication-efficient learning; poisoning-resilient private FL; এবং auditable governed systems। Anchor studies-এর capability map দেখায় যে প্রত্যেকে নির্দিষ্ট অংশে শক্তিশালী, কিন্তু exact integrated protocol কেউ evaluate করেনি।",
    "একটি paper accuracy-তে ভালো, আরেকটি cryptography-তে strong - তাই তাদের number সরাসরি compare করলে ভুল হবে। আমরা method, privacy unit, server visibility, transport, attack model এবং governance assumption দিয়ে synthesis করেছি।",
    ["FLDP: malicious-secure private aggregation", "FAMHE: biomedical multiparty HE", "DDFed and DP-BREM+: privacy plus poisoning", "PREAMBLE: private block-sparse aggregation"],
    ["Natural healthcare evidence", "Formal DP", "Protected aggregation", "Communication co-design", "Poisoning defence", "Distributed trust"],
    [
        ("আপনার তিনটি সবচেয়ে গুরুত্বপূর্ণ base paper কোন ধরনের?", "একটি encrypted/private aggregation direction, একটি privacy-aware or poisoning-resilient integrated method, এবং একটি communication-aware protected aggregation direction। Thesis তাদের capability combine করার gap থেকে motivated।"),
        ("Related paper-এর result reproduce করেছেন?", "Relevant recent baselines Fed-EF-SGD, Fed-EF-AMS এবং adapted SoteriaFL-SGD একই healthcare protocol-এ run করা হয়েছে। Published headline score replicate দাবি করা হয়নি, কারণ dataset ও privacy configuration আলাদা।"),
    ],
)

h2("2.3.1 Multi-Institutional Healthcare Federated Learning")
p("Natural multi-hospital studies FL-এর feasibility এবং institutional heterogeneity দেখিয়েছে। FLamby-এর মতো benchmark reproducible healthcare tasks দিয়েছে। কিন্তু অনেক কাজ formal patient-level DP, encrypted wire measurement, poisoning evaluation বা governance একসঙ্গে দেয় না। আমাদের dataset split natural hospital system ধরে এই evidence stream-এর সঙ্গে align করে।")

h2("2.3.2 Privacy-Enhanced Healthcare Learning")
p("এই stream DP, secure aggregation, multiparty HE এবং privacy attack analysis ব্যবহার করে। শক্তি হলো formal বা cryptographic confidentiality; দুর্বলতা হলো utility cost, key management, computational overhead এবং limited threat integration। Thesis DP এবং CKKS-এর responsibility আলাদা রেখে উভয় ব্যবহার করেছে।")

h2("2.3.3 Communication-Efficient and Encrypted Federated Learning")
p("Fed-EF family biased compression-এর জন্য error feedback analyse করে; SoteriaFL privacy-aware compression integrate করে; PREAMBLE block-sparse protected aggregation co-design করে। আমাদের adaptation record-level DP, healthcare weighting এবং dense CKKS transport-এর সঙ্গে same-protocol comparison দেয়। প্রধান gap ছিল actual serialized ciphertext saving।")

h2("2.3.4 Poisoning-Resilient and Privacy-Aware Federated Learning")
p("DDFed encrypted anomaly detection এবং DP-BREM+ privacy ও Byzantine resilience combine করার চেষ্টা করে। এরা compatibility gap আংশিক সমাধান করে, কিন্তু cryptographic path, adversary, dataset এবং trust assumptions ভিন্ন। তাই তাদের published scores contextual, direct baseline নয়।")

h2("2.3.5 Auditable and Governed Federated Systems")
p("Blockchain বা permissioned-ledger based FL works identity, provenance এবং immutable logging তুলে ধরে। কিন্তু ledger দিয়ে model correct, patient-private বা poisoning-resistant প্রমাণ হয় না। আমাদের design audit-কে signed checkpoint evidence হিসেবে ব্যবহার করে এবং technical control-এর সীমা আলাদা রাখে।")

section_lesson(
    "2.4 Comparative Analysis of Existing Approaches",
    "Approaches utility, formal privacy, update confidentiality, real communication saving, poisoning defence, governance এবং central trust assumption দিয়ে compare করা হয়েছে। এই matrix দেখায় কোনো method সব dimension-এ dominant নয়।",
    "Fair comparison মানে একই প্রশ্ন করা: data কী, privacy budget কী, server কী দেখে, network byte মাপা হয়েছে কি, malicious fraction কত, defence-এর extra trust কী। শুধু AUROC column দিয়ে ranking যথেষ্ট নয়।",
    ["Predictive performance", "Privacy and security guarantee", "Communication and computation", "Robustness and availability", "Domain suitability"],
    ["Direct numeric comparison only under aligned protocol", "Literature result establishes context, not superiority"],
    [
        ("আপনার benchmark fair কেন?", "Same preprocessing, hospital split, model family, round count, DP mechanism, CKKS transport, seed pairing এবং evaluation code রাখা হয়েছে; method-specific recurrence শুধু প্রয়োজনমতো বদলেছে।"),
        ("All-in-one framework কি automatically best?", "না। Integration নতুন overhead ও failure mode আনে। তাই component ablation, multi-metric utility, attacks এবং service failure test দরকার।"),
    ],
)

section_lesson(
    "2.5 Summary of Research Gaps",
    "Literature থেকে পাঁচটি gap পাওয়া যায়: patient-level privacy ও update confidentiality-এর joint evaluation কম; logical compression প্রায়ই ciphertext saving প্রমাণ করে না; encrypted aggregation এবং plaintext robustness incompatible; key release ও audit governance অনেক architecture-এ weak; এবং healthcare evidence প্রায়ই ছোট বা artificial split-এ সীমিত।",
    "এই gaps সরাসরি methodology তৈরি করেছে: DP plus CKKS, CAABC aligned support, validation/reputation after aggregation, authority recomputation plus witness ledger, এবং full natural cohort with paired statistics।",
    ["Gap 1 -> DP plus CKKS", "Gap 2 -> CAABC real byte accounting", "Gap 3 -> poisoning evaluation and governed release", "Gap 4 -> verification and witness audit", "Gap 5 -> natural full cohort and frozen protocol"],
    ["Research gap specific, evidence-backed এবং solvable scope-এ হওয়া উচিত।"],
    [
        ("সবচেয়ে গুরুত্বপূর্ণ research gap কোনটি?", "Ciphertext-aware efficiency এবং privacy-robustness compatibility gap। এগুলো ছাড়া sparse encrypted FL হয় network saving হারায়, নয়তো defence-এর জন্য confidentiality ছাড়তে হয়।"),
    ],
)

section_lesson(
    "2.6 Organization of the Chapter",
    "Chapter theory দিয়ে vocabulary তৈরি করে, related studies দিয়ে evidence synthesize করে, comparative analysis দিয়ে assumptions expose করে, এবং research gaps দিয়ে proposed methodology justify করে।",
    "Literature review শেষে পাঠক যেন বুঝতে পারে proposed components arbitrary নয়; প্রতিটি design decision একটি observed gap-এর উত্তর।",
    ["Theory -> evidence -> comparison -> gap -> design"],
    ["Roadmap section concise রাখবে।"],
    [
        ("Literature review থেকে আপনার সবচেয়ে বড় lesson কী?", "Privacy-preserving FL-এ individual mechanism-এর strength যোগ করলেই secure system হয় না; assumptions compatible হতে হবে এবং joint trade-off measure করতে হবে।"),
    ],
)

h2("Chapter 2 Rapid Revision")
key_box("90 সেকেন্ডে Chapter 2", [
    "FL raw data local রাখে, কিন্তু update privacy নিশ্চিত করে না।",
    "DP final release leakage bound করে; CKKS in-transit/in-use update confidentiality দেয়।",
    "Compression useful only when serialized ciphertext bytes কমে।",
    "Plaintext robust aggregation privacy path-এর সঙ্গে conflict করতে পারে।",
    "Audit evidence sequence ও authorization দেখায়, correctness নয়।",
    "Research gap থেকে CAABC, governed release এবং verification architecture এসেছে।",
])


# ---------------------------------------------------------------------------
# CHAPTER 3
# ---------------------------------------------------------------------------
chapter_break()
h1("Chapter 3 Proposed Methodology")
p("এই chapter defense-এর technical core। Board জানতে চাইবে system কীভাবে কাজ করে, privacy ও security assumptions কী, equations কী বোঝায়, data leakage কীভাবে আটকানো হয়েছে এবং experiment fair কি না। উত্তর দেওয়ার সময় component-এর নামের চেয়ে data flow এবং trust boundary পরিষ্কার করবে।")

section_lesson(
    "3.1 Introduction",
    "Methodology literature gap-কে executable architecture ও controlled experiment-এ রূপ দেয়। Study design-science artifact construction-এর সঙ্গে quantitative evaluation যুক্ত করে। Logical flow হলো conceptual framework, threat model, mathematical formulation, architecture, data procedure এবং societal responsibility।",
    "এখানে ‘আমরা কী বানিয়েছি’ এবং ‘কীভাবে যাচাই করেছি’ দুটো আলাদা কিন্তু connected। Results এই chapter-এ নয়; এখানে pre-defined protocol এবং measurement plan থাকবে।",
    ["Method before result", "Frozen configuration", "Claim boundaries documented before interpretation"],
    ["Artifact construction", "Controlled evaluation", "Reproducible configuration"],
    [
        ("Methodology এবং implementation-এর পার্থক্য কী?", "Methodology research design, assumptions, variables ও evaluation logic ব্যাখ্যা করে; implementation সেই design-এর executable realization। Thesis দুইটি traceable রেখেছে।"),
    ],
)

section_lesson(
    "3.2 Conceptual Framework",
    "Conceptual framework সাতটি layer-এ ভাবা যায়: local data custody; patient-level privacy; ciphertext-aligned compression; CKKS and authenticated transport; aggregate verification; validation-reputation-release governance; এবং persistent audit. Output হলো governed global model, raw patient data নয়।",
    "Framework-এর প্রতিটি arrow একটি trust transition। Hospital raw data দেখে, coordinator ciphertext দেখে, authority aggregate plaintext release boundary, validator model behaviour দেখে, witness event digest রাখে। Role separation একজন actor-এর ক্ষমতা সীমিত করে।",
    ["Data minimization", "Separation of duties", "Defence in depth", "Evidence-linked release"],
    ["Local data -> DP update -> CAABC -> CKKS -> verify -> validate -> audit"],
    [
        ("Conceptual framework-এর dependent variable কী?", "Joint outcome: predictive utility, privacy exposure, serialized communication, poisoning impact এবং accountable release behaviour। একটি single scalar objective দিয়ে সব quality ধরা হয়নি।"),
        ("কেন layered design?", "প্রতিটি layer একটি আলাদা failure mode address করে এবং assumption স্পষ্ট করে। Layered design component replace বা audit করাও সহজ করে।"),
    ],
)

section_lesson(
    "3.2.1 Research Design and Methodological Propositions",
    "Three propositions ছিল: integrated protected workflow feasible; CAABC-FL dense encrypted reference-এর তুলনায় real wire কমিয়ে AUROC non-inferior থাকতে পারে; এবং security controls registered attacks-এ evaluate করা যায় without universal guarantees। Co-primary gate: AUROC difference-এর 95 শতাংশ lower bound -0.015-এর উপরে এবং wire reduction lower bound 0.40-এর উপরে।",
    "দুটি gate একসঙ্গে pass করতে হয়েছে। শুধু 56 শতাংশ byte saving কিন্তু utility collapse হলে method সফল নয়; আবার AUROC same কিন্তু traffic না কমলে contribution নেই। Pre-specified margin researcher bias কমায়।",
    ["LCB95(mean AUROC difference) > -0.015", "LCB95(mean wire reduction) > 0.40", "AUPRC and security outcomes separately interpreted"],
    ["Non-inferiority is not equality or superiority", "Co-primary decision prevents cherry-picking"],
    [
        ("-0.015 margin কেন?", "এটি study protocol-এ pre-specified acceptable AUROC loss, communication benefit-এর বিনিময়ে clinically oriented tolerance হিসেবে ব্যবহৃত। Margin choice domain judgement; future clinical validation প্রয়োজন।"),
        ("Wire reduction 40 শতাংশ gate কেন?", "Cryptographic overhead বিবেচনায় improvement practically meaningful করতে substantial minimum target দরকার ছিল। Observed reduction প্রায় 56 শতাংশ, margin-এর উপরে।"),
    ],
)

section_lesson(
    "3.2.2 Threat Model and Trust Assumptions",
    "Hospitals authenticated হলেও subset malicious হতে পারে। Coordinator honest-but-curious এবং aggregation protocol follow করে বলে confidentiality analysis ধরে। Network attacker eavesdrop, tamper, replay বা impersonate করতে পারে। Authority policy-constrained কিন্তু main service-এ trusted plaintext boundary। Witness outage বা single compromise considered; majority collusion excluded।",
    "Threat model security claim-এর fence। আমরা যে attacker test করিনি তার বিরুদ্ধে protection দাবি করি না। One-round attacks-এ adaptive-norm, ALIE, backdoor, Gaussian, label-flip, model-replacement ও sign-flip অন্তর্ভুক্ত ছিল। Longitudinal test repeated rounds-এ governance effect দেখে।",
    ["Protected assets: patient record, update, key, model integrity, audit order", "Trust assumptions: authenticated consortium, non-compromised authority, witness quorum", "Exclusions explicit"],
    ["Honest", "Honest-but-curious", "Byzantine", "Compromised service"],
    [
        ("Coordinator malicious হলে?", "Current confidentiality analysis protocol-following honest-but-curious coordinator ধরে। Aggregate substitution independent recomputation detect করতে পারে, কিন্তু arbitrary malicious coordinator বা collusion-এর formal security proof নেই।"),
        ("Authority compromise হলে?", "Main TenSEAL path-এর secret key authority-র কাছে, তাই authority compromise excluded high-impact risk। Networked threshold custodians future work-এর highest priority।"),
    ],
)

section_lesson(
    "3.3.1 Collaborative Learning Setting",
    "Institution k-এর dataset D_k={(x_ki,y_ki)} এবং local risk F_k(theta)=1/n_k sum loss। Global objective F(theta)=sum alpha_k F_k(theta), যেখানে alpha_k=n_k/total n। Source-system partition non-IID হতে পারে এবং raw records coordinator-এ pool করা হয় না।",
    "Equation-এর অর্থ: প্রতিটি hospital নিজের average loss কমায়; global model hospital sample count অনুযায়ী তাদের loss combine করে। Weighting বড় hospital-এর বেশি evidence reflect করে, তবে fairness perspective-এ equal-site weighting ভিন্ন design হতে পারে।",
    ["p=280 input features", "Model parameter dimension d=18,049", "Train only local; validation/test updates-এ ব্যবহৃত নয়"],
    ["Dataset ownership", "Weighted empirical risk", "Non-IID distributions"],
    [
        ("Sample-size weighting bias তৈরি করতে পারে?", "হ্যাঁ, বড় site বেশি influence পায়। এটি pooled empirical-risk objective-এর standard choice, কিন্তু institutional fairness বা equal-site benefit আলাদা objective; future work-এ evaluate করা উচিত।"),
    ],
)

section_lesson(
    "3.3.2 Patient-Level Privacy Constraint",
    "Neighbouring datasets D ও D prime একজন patient-এর সম্পূর্ণ contribution-এ ভিন্ন। Mechanism M-এর জন্য Pr[M(D) in O] <= exp(epsilon) Pr[M(D prime) in O] + delta। Per-sample clipping C=1.0, Gaussian noise multiplier 1.2, Poisson sampling এবং RDP accounting ব্যবহার হয়েছে।",
    "Definition বলে একজন patient আছে বা নেই - এই পরিবর্তনে output distribution খুব বেশি বদলবে না। Clipping sensitivity bound করে; noise indistinguishability বাড়ায়; accountant 20 round-এর local steps compose করে final ε=1.935 দেয়।",
    ["Privacy unit = patient", "Final budget = (1.935, 10^-5)", "CAABC support privatized update থেকে, তাই post-processing"],
    ["DP protects influence, not correctness", "Privacy budget whole training over compose হয়"],
    [
        ("Gradient clipping নিজেই DP কি?", "না। Clipping sensitivity bound করে, কিন্তু calibrated random noise ও accountant ছাড়া DP guarantee হয় না।"),
        ("DP noise কোথায় যোগ হয়েছে?", "Hospital-side per-sample DP-SGD training-এর সময়, encryption ও compression-এর আগে। ফলে coordinator-এর কাছে যাওয়া update ইতিমধ্যে privatized।"),
    ],
)

section_lesson(
    "3.3.3 Ciphertext-Aligned Communication Constraint",
    "Privatized update plus residual u_k,t block-এ partition হয়। B=256 coordinate, CKKS slot L=4096, target c=2 ciphertext chunks। Common support S_t এমন block বেছে নেয় যাতে selected coordinates cL-এর মধ্যে থাকে। Projection transmitted হয় এবং omitted part residual r_k,t+1-এ যায়।",
    "18,049 parameter dense vector কয়েকটি CKKS chunk চায়। CAABC useful block common support-এ রাখে, ফলে সব hospital একই slot map অনুসরণ করে এবং aggregate সরাসরি যোগ করা যায়। Refresh প্রতি পাঁচ round-এ support update করে।",
    ["Block size 256", "Two update ciphertext chunks", "Refresh interval 5", "Error feedback enabled"],
    ["Same support prevents coordinate misalignment", "Support scores already-private updates থেকে aggregate", "Wire objective includes refresh overhead"],
    [
        ("Shared support privacy risk নয়?", "Support কিছু aggregate structural information প্রকাশ করতে পারে, কিন্তু selection privatized update-এর post-processing। Formal DP budget বাড়ে না; তবুও support leakage future analysis-এর topic।"),
        ("Block contiguous কেন?", "Contiguous blocks deterministic packing, index metadata কমানো এবং CKKS chunk alignment সহজ করে। Arbitrary coordinate set বেশি metadata ও alignment complexity আনতে পারে।"),
    ],
)

section_lesson(
    "3.3.4 Authenticated and Verified Encrypted Aggregation",
    "Hospital packed update public CKKS context দিয়ে encrypt করে এবং identity, round, model version, privacy metadata, sample count, support ID, nonce, timestamp ও digestসহ signed envelope পাঠায়। Coordinator admissible envelopes weighted homomorphic sum করে। Authority same envelope independently verify ও recompute করে digest match হলে aggregate decrypt করে।",
    "Signature বলে message কে পাঠিয়েছে এবং বদলেছে কি না; nonce replay আটকায়; model/support/context binding wrong-round update আটকায়; encryption content লুকায়; recomputation coordinator omission বা reweighting detect করে।",
    ["Coordinator holds no TenSEAL secret key", "Aggregate digest match required", "Individual ciphertext intended flow-এ decrypt হয় না"],
    ["mTLS channel", "Digital signature", "Nonce freshness", "Deterministic recomputation"],
    [
        ("mTLS থাকলে signature কেন?", "mTLS connection authenticate করে, কিন্তু signed envelope connection-এর বাইরে persistent non-repudiable evidence ও audit binding দেয়। Stored artifact পরে verify করা যায়।"),
        ("Homomorphic average কীভাবে হয়?", "Coordinator ciphertext addition এবং plaintext scalar multiplication দিয়ে weighted sum করে। CKKS approximate arithmetic-এর কারণে decrypted value expected weighted average-এর খুব কাছাকাছি, exact bitwise equal নয়।"),
        ("Independent recomputation-এর সীমা কী?", "Authority original envelopes পেলে omission/substitution detect করে, কিন্তু এটি zero-knowledge proof নয় এবং authority নিজে trusted।"),
    ],
)

section_lesson(
    "3.3.5 Governed Model Release and Accountability",
    "Decrypted aggregate দিয়ে candidate model তৈরি হলেও সেটি automatic global model হয় না। Validator quorum performance ও integrity rule check করে। Pass হলে model registry-তে signed version commit; fail হলে previous model rollback। Participant reputation update এবং severe condition-এ quarantine হয়। Audit witnesses checkpoint record করে।",
    "এটি prevention-এর বদলে controlled release। Poisoned update aggregate-এ ঢুকতে পারে, কিন্তু harmful candidate validation-এ ধরা পড়লে publish হয় না। Cost হলো false rollback হলে useful learning বন্ধ হতে পারে।",
    ["Release condition separates compute from publish", "Rollback rate availability metric", "Audit stores pseudonymized/minimized evidence"],
    ["Validation quorum", "Model versioning", "Rollback", "Reputation", "Witness receipts"],
    [
        ("Validation data কোথায়?", "Training data থেকে disjoint validation partition-এ threshold ও release behaviour assess করা হয়েছে; test set tuning বা gate decision-এ ব্যবহার করা হয়নি।"),
        ("Backdoor validation-এ ধরা না পড়লে?", "Possible। Current gate registered attacks ও validation signals-এর জন্য empirical control, formal guarantee নয়। Adaptive stealthy backdoor future risk।"),
    ],
)

section_lesson(
    "3.3.6 Joint Research Objective",
    "Joint objective prediction loss minimise করার পাশাপাশি privacy budget, ciphertext packing, wire cost, authentication, aggregate verification এবং release policy satisfy করা। এটি true multi-objective engineering problem; একটি single best point নেই।",
    "একটি method accuracy সামান্য কমিয়ে network cost অনেক কমাতে পারে। আবার aggressive defence poisoning কমিয়ে rollback বাড়াতে পারে। তাই Pareto-style trade-off এবং pre-defined decision gate ব্যবহার করা হয়েছে।",
    ["Utility", "Privacy", "Communication", "Integrity", "Availability"],
    ["Primary conclusion joint gates-এ", "Secondary metrics negative হলেও report করা হয়েছে"],
    [
        ("Single objective function কেন দেননি?", "Clinical utility, privacy loss, bytes এবং attack impact different units ও stakeholder value বহন করে। Arbitrary scalar weight-এর বদলে constraints, gates এবং multi-metric reporting বেশি transparent।"),
    ],
)

section_lesson(
    "3.4 System Architecture",
    "Architecture modular cross-silo client-server। Hospital service local training ও encryption করে; coordinator round orchestration ও public-key aggregation; authority envelope reverification, decryption policy ও candidate release; validators model check; registry signed version রাখে; witnesses external audit checkpoint রাখে।",
    "End-to-end flow: global model receive -> local DP train -> CAABC compress -> CKKS encrypt -> sign and submit -> verify and aggregate -> authority recompute and decrypt -> validate -> commit or rollback -> audit -> next round।",
    ["Seven logical layers", "Least privilege roles", "Persistent state and failure handling"],
    ["L1 data acquisition", "L2 privacy and crypto", "L3 local training and compression", "L4 trust management", "L5 aggregation", "L6 distribution", "L7 governed inference"],
    [
        ("Coordinator compromise হলে raw data পাবে?", "Raw record কখনও পায় না এবং intended path-এ secret CKKS context নেই। তবে metadata, ciphertext size, timing ও aggregate protocol behaviour side information হতে পারে।"),
        ("Architecture-এর single point of trust কোথায়?", "Main service-এ research authority, কারণ complete TenSEAL secret context তার কাছে। Witness ও recomputation coordinator risk কমায়, authority risk eliminate করে না।"),
    ],
)

h2("3.4.1 Algorithms and Model Explanation")
h3("Prediction model and local objective")
p("Model 280 input feature থেকে 64 hidden unit এবং একটি binary output-এর compact MLP; মোট 18,049 trainable parameter। Binary cross-entropy local loss এবং 20 communication round, one local epoch, batch size 128, learning rate 0.03 ব্যবহার হয়েছে। Compact model controlled security experiments এবং CKKS cost manageable রাখে।")
qa("Deep network ব্যবহার করেননি কেন?", "Primary goal architecture ও privacy-communication-security trade-off isolate করা। Compact MLP reproducible এবং expensive attack/CKKS matrix feasible করেছে। Larger models-এ scalability future work।")

h3("Patient-level differentially private training")
p("প্রতি patient gradient norm C=1.0-তে clip করা হয়, aggregate gradient-এ Gaussian noise যোগ হয় এবং secure Poisson sampling-based accountant final privacy loss track করে। একই frozen mechanism dense and CAABC private methods-এ ব্যবহার করায় compression effect আলাদা করে দেখা যায়।")
qa("DP comparison fair কীভাবে?", "Private baselines এবং CAABC একই clipping, noise multiplier, sampling, rounds ও accountant ব্যবহার করেছে। তাই utility difference মূলত compression বা optimizer configuration থেকে, privacy budget difference থেকে নয়।")

h3("CAABC block selection and error feedback")
p("প্রথমে effective update u = DP update + residual। প্রতিটি 256-coordinate block-এর score privatized energy থেকে তৈরি হয়। Refresh round-এ encrypted aggregate score থেকে common top blocks নির্ধারিত হয়। Selected blocks pack ও encrypt করা হয়; omitted components residual-এ থাকে। পরের round-এ residual নতুন private update-এর সঙ্গে যোগ হয়।")
qa("Adaptive বলতে কী adapt করে?", "Shared support privatized aggregate block importance অনুযায়ী periodic refresh হয়। Static CAABC fixed support রাখে; adaptive CAABC changing optimization signal অনুসরণ করে।")

h3("Encrypted aggregation and governed update")
p("Selected vector CKKS public context-এ encrypt হয়। Accepted ciphertext weighted sum করে authority-verified aggregate তৈরি হয়। Decrypted aggregate full coordinate space-এ scatter করে candidate global model update হয়। Validation pass না করলে previous signed model থাকে।")
qa("Encrypted domain-এ poisoning detect করেন?", "Main path individual encrypted anomaly score দিয়ে formal robust aggregation করে না। Independent recomputation protocol integrity দেখে; model-level validation ও reputation post-aggregation poisoning impact control করে।")

h3("Parameter settings")
two_col_table(
    ["Component", "Frozen setting", "কারণ"],
    [
        ("Clinical representation", "280 features, 24-hour history", "Matched patient-level landmark representation"),
        ("Model", "280-64-1 MLP, 18,049 parameters", "Controlled utility and feasible crypto/attack cost"),
        ("Training", "20 rounds, 1 local epoch, batch 128, LR 0.03", "Consistent cross-method protocol"),
        ("DP", "C=1.0, sigma=1.2, delta=10^-5", "Recorded patient-level accountant"),
        ("CKKS", "N=8192, 60-40-40-60, scale 2^40", "Packed approximate aggregation"),
        ("CAABC", "Block 256, 2 chunks, refresh 5", "Meaningful wire reduction with utility retention"),
    ],
    [1.3, 2.2, 3.0],
)

h3("Complexity and communication analysis")
p("Dense encryption cost parameter dimension ও ciphertext chunk count-এর সঙ্গে বাড়ে। CAABC score computation O(d), block selection roughly O(q log k), encryption selected coordinates-এর chunk count অনুযায়ী। Real benefit plaintext nonzero count নয়, serialized objects-এর B_method/B_dense ratio দিয়ে measure করা হয়েছে। Metadata ও score refresh overhead অন্তর্ভুক্ত।")
qa("Compression overhead saving খেয়ে ফেলেনি?", "Adaptive refresh extra message তৈরি করেছে, তাই static version-এর চেয়ে সামান্য বেশি traffic; তবুও measured final reduction 55.998 শতাংশ ছিল এবং target 40 শতাংশের বেশি।")

section_lesson(
    "3.4.2 Data Flow Diagram",
    "DFD raw clinical record-এর boundary থেকে governed model release পর্যন্ত data object, process এবং store দেখায়। Raw data hospital boundary ছাড়ে না; transmitted object হলো DP-sanitized compressed ciphertext plus signed metadata; persistent artifacts হলো model version, envelope digest, consent/purpose record এবং audit checkpoint।",
    "Diagram ব্যাখ্যা করতে actor নয়, data object follow করবে। কোন পর্যায়ে plaintext, privatized vector, ciphertext, aggregate এবং model আছে তা বললে trust boundary পরিষ্কার হয়।",
    ["Raw data local", "DP before encryption", "Coordinator sees ciphertext", "Authority sees verified aggregate", "Audit sees digest, not patient identity"],
    ["Training flow", "Release flow", "Inference governance flow"],
    [
        ("Data leakage সবচেয়ে sensitive কোথায়?", "Hospital local training environment ও authority decryption boundary। Coordinator individual plaintext পায় না; audit log-এ raw patient pseudonymও রাখা হয়নি।"),
    ],
)

section_lesson(
    "3.5.1 Data Collection Methods",
    "PhysioNet Computing in Cardiology Challenge 2019-এর complete 40,336-patient cohort ব্যবহার করা হয়েছে। দুই source hospital system natural institutional partition হিসেবে রাখা হয়। 24-hour history থেকে 280-feature matched-landmark representation তৈরি করা হয়েছে। Public licensed retrospective data ব্যবহৃত হয়েছে; নতুন patient recruitment হয়নি।",
    "Matched landmark মানে সব patient-এর জন্য comparable prediction time ও history window তৈরি করা, যাতে future information accidentally feature-এ না ঢোকে। Patient-level split একই patient-এর row train ও test উভয়ে যাওয়ার leakage আটকায়।",
    ["40,336 patients", "Two natural hospitals", "No synthetic client fragmentation", "Patient-level split"],
    ["Data provenance", "Inclusion logic", "Missingness handling", "Leakage control"],
    [
        ("একটি dataset যথেষ্ট?", "Internal evidence-এর জন্য full natural cohort strong, কিন্তু external validity-এর জন্য যথেষ্ট নয়। Independent cohort, more hospitals এবং prospective silent validation future work।"),
        ("Why PhysioNet 2019?", "Public multi-centre sepsis data, temporal clinical variables এবং natural hospital systems আছে; reproducible cross-silo evaluation সম্ভব।"),
    ],
)

h2("3.5.1.1 Clinical Data Source and Relevance")
p("Endpoint early sepsis risk modelling, যেখানে positive class তুলনামূলক কম। তাই AUROC-এর পাশাপাশি AUPRC, sensitivity, precision ও calibration জরুরি। Dataset research benchmark; result clinical deployment evidence নয়।")

h2("3.5.1.2 Patient-Level Sampling and Leakage Control")
p("সব measurement একই patient identifier অনুযায়ী group করে train, validation ও test ভাগ হয়েছে। Test labels hyperparameter, threshold, compression support বা release gate নির্বাচন করতে ব্যবহার হয়নি। Validation threshold pooled validation partition থেকে একবার select করে test evaluation-এর আগে freeze করা হয়।")
qa("Data leakage কীভাবে আটকেছেন?", "Patient-level partition, training-only preprocessing fit, fixed landmark, validation-only threshold selection এবং test set final evaluation-এর জন্য রেখে।")

h2("3.5.1.3 Matched-Landmark Preprocessing")
p("Clinical variables temporal window-এ summarize, missingness indicator ও derived statisticsসহ fixed 280-dimensional patient vector তৈরি হয়েছে। একই preprocessing contract সব hospital ও baseline-এ প্রয়োগ হওয়ায় feature-space consistency বজায় থাকে।")

h2("3.5.1.4 Experimental and System Measurements")
p("Scientific experiments utility, DP, CKKS, compression, attacks ও statistics মাপে। Service profile mTLS, signatures, persistence, authority reverification, witness quorum ও failure recovery মাপে। দুই profile complementary; complete 40,336-patient run এক physical production cluster দিয়ে চালানো হয়নি।")

section_lesson(
    "3.5.2 Data Analysis Techniques",
    "Predictive metrics, calibration, serialized communication, privacy attacks, reconstruction, poisoning, cryptographic error, latency এবং throughput measure করা হয়েছে। Same-seed paired difference, 95 শতাংশ confidence interval, paired t-test, Wilcoxon sensitivity এবং Holm multiplicity correction ব্যবহার হয়েছে।",
    "Mean score শুধু effect size বলে না। Confidence interval estimate কত uncertain তা দেখায়। Paired design same experimental randomness share করে। Holm family-wise false positive কমায়। Primary conclusion p-value নয়, pre-specified gate-এর উপর।",
    ["Ten seeds core utility", "Five locked seeds costly security/SOTA", "Attack protocol frozen", "Negative results retained"],
    ["Descriptive vs confirmatory", "Direct vs contextual comparison", "Effect size and interval"],
    [
        ("মাত্র পাঁচ seed কেন?", "Security, reference-model membership inference এবং 300-step reconstruction computationally costly ছিল। Paired locked design uncertainty estimate দেয়, কিন্তু broad distribution characterize করে না; limitation stated।"),
        ("t-test পাঁচ seed-এ reliable?", "Small n-এ normality assumption fragile, তাই confidence interval-এর সঙ্গে Wilcoxon sensitivity ও pre-specified gates ব্যবহৃত। Result overgeneralize করা হয়নি।"),
    ],
)

h2("3.5.2.1 Comparative Experimental Matrix")
p("Matrix-এ non-private optimizers, dense DP, dense DP plus CKKS, logical compressors, CKKS-compatible block methods, CAABC variants, Fed-EF, SoteriaFL, robust aggregators এবং attack conditions রয়েছে। একই protocol-এ fairness বজায় রেখে প্রতিটি method-এর compatible setting clearly recorded।")

h2("3.5.2.2 Predictive and Calibration Analysis")
p("Accuracy, precision, sensitivity, specificity, F1, AUROC, AUPRC, balanced accuracy, MCC এবং Brier score ব্যবহার হয়েছে। Imbalanced sepsis endpoint-এ AUPRC positive-class quality এবং Brier probability calibration error দেখায়।")

h2("3.5.2.3 Communication Privacy and Security Analysis")
p("Communication serialized CKKS bytes; MIA loss-based, LiRA ও RMIA; reconstruction 300 optimization step; poisoning seven attack families at 40 শতাংশ malicious clients; longitudinal backdoor five rounds; governance outcomes rollback, quarantine ও forged-aggregate rejection।")

h2("3.5.2.4 Statistical Inference and Reproducibility")
p("Seeds, hashes, dependency versions, model configuration, dataset pipeline এবং artifact paths record করা হয়েছে। Reproducibility মানে identical hardware-এ bitwise result বাধ্যতামূলক নয়; protocol ও uncertainty reproduce করা।")

section_lesson(
    "3.6.1 Societal Impacts of Engineering Solutions",
    "Potential benefit হলো broader evidence থেকে model, institutional autonomy এবং auditability। Risk হলো biased model, false reassurance, exclusion of low-resource hospitals, governance burden এবং automation bias। System human clinical responsibility replace করে না।",
    "Privacy technology থাকলেও unequal data বা access থাকলে social harm হতে পারে। Institution-এর benefit, patient consent এবং accountability jointly evaluate করা দরকার।",
    ["Benefit distribution", "Clinical harm from false negative/positive", "Human oversight"],
    ["Research model != medical device"],
    [
        ("Societal benefit কীভাবে measure করেছেন?", "Direct user study বা clinical outcome measure করা হয়নি। Societal discussion potential impact ও risk assessment; empirical claims predictive এবং system metrics পর্যন্ত সীমিত।"),
    ],
)

section_lesson(
    "3.6.2 Environment and Sustainability Considerations",
    "FL raw data duplication কমাতে পারে, কিন্তু repeated local compute ও cryptography energy cost বাড়ায়। CAABC encrypted traffic ও potentially transmission energy কমায়। Compact model এবং open-source CPU execution resource need সীমিত করেছে।",
    "56 শতাংশ upload saving energy saving-এর direct measurement নয়। Network energy topology-dependent; তাই environmental benefit সম্ভাব্য, measured carbon claim নয়।",
    ["Compute-communication trade-off", "No measured carbon footprint"],
    ["Sustainability claim evidence-bounded"],
    [
        ("CAABC green technology কি?", "এটি encrypted network traffic কমায়, যা energy reduction-এর সম্ভাবনা দেয়; কিন্তু actual power বা carbon measure করা হয়নি, তাই green superiority claim করা যাবে না।"),
    ],
)

section_lesson(
    "3.6.3 Ethical and Professional Responsibility",
    "Ethical principles হলো data minimization, purpose limitation, consent binding, least privilege, transparency, reproducibility, secure key management, honest limitation reporting এবং human accountability। Public data হলেও re-identification ও misuse risk বিবেচনা করা হয়েছে।",
    "Professional responsibility মানে negative result লুকানো নয়। AUPRC loss, poisoning vulnerability, centralized authority এবং no real TDX/SNP experiment স্পষ্টভাবে লেখা research integrity-এর অংশ।",
    ["Do not overclaim", "No raw identifiers in audit", "Purpose-bound inference", "External audit before deployment"],
    ["Privacy by design", "Responsible disclosure", "Clinical non-deployment boundary"],
    [
        ("আপনার thesis-এর ethical strength কী?", "Patient-level privacy formalization, local custody, minimized audit metadata, consent-purpose checks এবং evidence-bounded reporting।"),
        ("Ethical limitation?", "No patient, clinician বা data-protection officer user study; fairness subgroup analysis ও prospective workflow impact future work।"),
    ],
)

h2("Chapter 3 Rapid Revision")
key_box("দুই মিনিটে Methodology", [
    "Data: 40,336 patients, two natural hospitals, 280 features, patient-level split।",
    "Model: 280-64-1 MLP, 18,049 parameters, 20 rounds।",
    "DP: C=1.0, sigma=1.2, delta=10^-5, final epsilon=1.935।",
    "CAABC: 256-coordinate blocks, 2 chunks, refresh every 5 rounds, error feedback।",
    "CKKS: N=8192, scale 2^40, 60-40-40-60 chain।",
    "Primary gates: AUROC LCB > -0.015 and wire reduction LCB > 40%。",
    "Security: signed envelopes, nonce/version binding, authority recomputation, validation, rollback, reputation, witness audit।",
    "Boundary: main service centrally keyed; OpenFHE 3-of-3 experiment separate।",
])


# ---------------------------------------------------------------------------
# CHAPTER 4
# ---------------------------------------------------------------------------
chapter_break()
h1("Chapter 4 Results and Discussion")
p("এই chapter-এ board সবচেয়ে বেশি evidence চাইবে। Number বলার পাশাপাশি ‘relative to what’, ‘under which protocol’, ‘how uncertain’, এবং ‘what cannot be concluded’ বলতে হবে। Strong defense answer positive ও adverse ফলাফল দুটোই স্বীকার করে।")

section_lesson(
    "4.1 Introduction",
    "Results চার প্রশ্নের উত্তর দেয়: privacy ও encryption-এর পরে কত utility থাকে; CAABC real encrypted traffic কমায় কি না; privacy এবং poisoning attack-এ behaviour কী; এবং protected service scale করলে system cost কী। Quantitative ও qualitative evidence পরে literature এবং objectives-এর সঙ্গে connect করা হয়েছে।",
    "Chapter 4-কে result dump ভাববে না। প্রতিটি table একটি research question answer করে এবং discussion claim-এর strength নির্ধারণ করে।",
    ["Utility", "Communication", "Privacy exposure", "Poisoning", "Systems and governance"],
    ["Result -> interpretation -> limitation"],
    [
        ("Primary result কোনটি?", "Dense DP-FedAvg-CKKS-এর তুলনায় CAABC-FL-এর joint AUROC non-inferiority এবং actual encrypted-wire reduction gate pass করা।"),
    ],
)

section_lesson(
    "4.2.1 Performance Metrics",
    "Classification metrics confusion matrix থেকে আসে। Accuracy=(TP+TN)/all, Precision=TP/(TP+FP), Sensitivity=TP/(TP+FN), Specificity=TN/(TN+FP), F1 precision ও sensitivity-এর harmonic mean। AUROC threshold-independent ranking, AUPRC imbalanced positive class-এর জন্য informative, Brier probability error measure।",
    "Sepsis positive case কম। সব patient negative বললেও accuracy misleadingly high হতে পারে। তাই sensitivity missed positive, precision false alarm burden, AUPRC positive ranking এবং Brier probability quality দেখায়।",
    ["AUROC high মানে calibration ভালো এমন নয়।", "AUPRC prevalence-sensitive।", "Threshold validation set-এ select এবং test-এর আগে freeze।"],
    ["Balanced accuracy=(sensitivity+specificity)/2", "MCC imbalance-aware correlation", "R_wire=1-B_method/B_dense", "Reconstruction RMSE larger means poorer attack reconstruction"],
    [
        ("Accuracy 0.83 কিন্তু F1 0.29 কেন?", "Class imbalance-এর কারণে true negatives accuracy dominate করে। Positive prediction quality দুর্বল হতে পারে; তাই F1 ও AUPRC দরকার।"),
        ("AUROC এবং AUPRC কেন দুটো?", "AUROC ranking across both classes; AUPRC positive class retrieval quality এবং imbalance-এর প্রতি বেশি sensitive। একই AUROC হলেও AUPRC কমতে পারে।"),
        ("Wire reduction কীভাবে হিসাব?", "1 minus method-এর serialized encrypted upload divided by dense encrypted reference upload। Ciphertext object serialize করে মাপা হয়েছে।"),
    ],
)

h2("4.2.2 Result Tables and Graphical Representations")
h3("Predictive Performance")
p("Ten paired seeds-এর mean results নিচে। Non-private FedAvg highest utility region-এ, কিন্তু formal DP বা update confidentiality নেই। Fair protected reference হলো dense DP-FedAvg-CKKS।")
two_col_table(
    ["Method", "Accuracy", "AUROC", "AUPRC", "Sensitivity", "Specificity", "Precision", "F1", "Brier"],
    [
        ("FedAvg", "0.821", "0.830", "0.471", "-", "-", "-", "-", "-"),
        ("DP-FedAvg", "0.832", "0.732", "0.394", "0.458", "0.861", "0.224", "0.293", "0.064"),
        ("DP-FedAvg-CKKS", "0.833", "0.731", "0.394", "0.457", "0.862", "0.228", "0.296", "0.064"),
        ("CAABC-static", "0.815", "0.732", "0.390", "0.476", "0.841", "0.210", "0.282", "0.065"),
        ("CAABC-FL", "0.758", "0.731", "0.384", "0.534", "0.776", "0.173", "0.253", "0.066"),
    ],
)
p("CAABC-FL AUROC dense encrypted baseline-এর সমান অঞ্চলে, কিন্তু operating point বদলেছে: sensitivity 0.457 থেকে 0.534 হয়েছে, specificity 0.862 থেকে 0.776 এবং F1 0.296 থেকে 0.253 হয়েছে। অর্থাৎ বেশি positive case detect হলেও false positive burden বেড়েছে।")
qa("CAABC-FL accuracy কম, তাহলে successful কেন?", "Primary utility criterion accuracy নয়, pre-specified AUROC non-inferiority; এটি pass করেছে এবং 56 শতাংশ encrypted traffic কমিয়েছে। তবে lower accuracy, precision, F1 এবং AUPRC adverse trade-off হিসেবে স্বীকার করা হয়েছে।")
qa("DP-তে utility কেন কমল?", "Per-sample clipping ও Gaussian noise optimization signal perturb করে। DP-FedAvg থেকে DP-FedAvg-CKKS প্রায় একই হওয়ায় major loss CKKS নয়, DP mechanism থেকে এসেছে।")

h3("Encrypted Communication Efficiency")
two_col_table(
    ["Method", "Encrypted upload per round", "Actual reduction", "Interpretation"],
    [
        ("Dense DP-FedAvg-CKKS", "3.314 MB", "0%", "Protected reference"),
        ("Top-k -> dense CKKS", "3.314 MB", "0%", "Logical sparse, wire dense"),
        ("EF-Top-k -> dense CKKS", "3.314 MB", "0%", "Error feedback, no ciphertext saving"),
        ("Fixed-block CKKS", "1.326 MB", "about 60%", "Common aligned static blocks"),
        ("CAABC-static", "1.359 MB", "about 59%", "Fixed learned support"),
        ("CAABC-FL", "1.458 MB", "55.998%", "Adaptive refresh overhead included"),
    ],
    [1.55, 1.55, 1.2, 2.2],
)
p("CAABC-FL client-specific sparse indices পাঠায় না; common support-এর selected block fewer CKKS chunks-এ pack হয়। Refresh score traffic count করেও 55.998 শতাংশ reduction হয়েছে। এটাই algorithmic contribution-এর সবচেয়ে direct evidence।")
qa("Fixed-block বেশি byte save করে, CAABC কেন?", "Fixed-block প্রায় 60 শতাংশ save করলেও support task dynamics adapt করে না। CAABC slightly lower 56 শতাংশ saving দিয়ে changing aggregate block importance follow করে এবং registered utility gate pass করেছে। Contribution adaptability plus measured transport, শুধু maximum compression নয়।")
qa("MB না MiB?", "Scientific table decimal MB convention ব্যবহার করেছে; service scalability আলাদা করে MiB report করেছে। Defense-এ table-এর unit অনুসরণ করবে এবং unit mix করবে না।")

h3("Poisoning Behaviour")
p("Confirmatory one-round matrix-এ সাত attack, ছয় defence/aggregator এবং five seeds; total 870 evaluated conditions। Severe point 40 শতাংশ malicious clients। Weighted encrypted-compatible mean model-replacement attack-এ AUROC delta -0.393 [ -0.421, -0.366 ] এবং backdoor ASR প্রায় +0.888। Plaintext FLTrust ও Multi-Krum কিছু attack-এ শক্তিশালী, কিন্তু individual geometry বা trusted root data চায়।")
p("Five-round longitudinal backdoor-এ validation, rollback ও reputation ব্যবহারে final AUROC delta -0.0293, ASR change +0.0046 যার interval zero cross করে, rollback rate 0.60 এবং malicious-client quarantine recall 1.00। Protection attack impact কমিয়েছে, কিন্তু 60 শতাংশ candidate round reject হওয়ায় availability এবং utility cost দেখিয়েছে।")
qa("Framework poisoning-proof?", "না। Ordinary encrypted weighted mean strong one-round model replacement-এ vulnerable। Governance registered longitudinal attack-এর release impact সীমিত করেছে, কিন্তু formal Byzantine guarantee নয়।")
qa("40 শতাংশ malicious realistic?", "এটি stress-test severity, সব deployment-এর prevalence claim নয়। বিভিন্ন fraction-এর মধ্যে 40 শতাংশ worst registered condition হিসেবে robustness boundary দেখিয়েছে।")
qa("ASR কী?", "Attack Success Rate - backdoor trigger দিলে attacker যে target output চায় তা model কতবার দেয়। Clean AUROC এবং ASR দুটো একসঙ্গে দেখতে হয়।")

h3("Membership Inference and Gradient Reconstruction")
two_col_table(
    ["Privacy test", "FedAvg/plaintext", "Dense DP-CKKS", "CAABC-FL", "Interpretation"],
    [
        ("Loss-based MIA AUC", "0.5239", "0.5050", "0.5035", "DP methods near chance"),
        ("RMIA AUC", "0.5492", "0.5067", "0.5063", "Reference-model attack weakened"),
        ("LiRA online AUC", "higher than chance", "near chance", "0.4963", "Attack-specific diagnostic"),
        ("300-step reconstruction RMSE", "0.6844", "2.5467", "2.7632", "Higher RMSE means poorer reconstruction"),
    ],
    [1.4, 1.35, 1.35, 1.25, 2.0],
)
p("Powered loss-based experiment প্রতি seed 720 member ও 720 non-member ব্যবহার করেছে। CAABC-FL confidence interval [0.4929, 0.5142], chance 0.5-এর আশেপাশে। DP plus CAABC reconstruction RMSE plaintext-এর তুলনায় 2.079 বেড়েছে। কিন্তু attack fail মানে privacy proof নয়; formal statement accountant-এর (ε,δ) থেকে।")
qa("Attack AUC 0.5-এর নিচে হলে কি perfect privacy?", "না। Attacker score direction reverse করতে পারে, confidence interval এবং attack design matter করে। Near 0.5 শুধু registered attacker discriminate করতে পারেনি; formal privacy DP guarantee থেকে।")
qa("CAABC কি DP-এর চেয়ে extra privacy দেয়?", "Reconstruction worst-case RMSE সামান্য বেশি, কিন্তু compression alone reliable privacy mechanism নয়। Formal budget CAABC-এর জন্য আলাদা stronger হয়নি, কারণ এটি DP update-এর post-processing।")
qa("Reconstruction RMSE বড় কেন ভালো?", "এই metric attacker-এর recovered feature এবং true feature-এর distance। বড় error মানে reconstruction কম accurate। Prediction RMSE-এর মতো এখানে lower ভালো নয়; attacker perspective-এ lower ভালো, defender perspective-এ higher।")

section_lesson(
    "4.2.3 Statistical Analysis",
    "Core comparisons paired by seed। Two-sided 95 শতাংশ confidence interval effect estimation; paired t-test plus Wilcoxon sensitivity; Holm correction endpoint family multiplicity control। Primary decision two gates। AUROC mean difference +0.00017, CI [-0.00212,+0.00247], lower bound -0.015 margin-এর উপরে। Wire reduction 55.998 শতাংশ, CI [55.994,56.002], lower bound 40 শতাংশ target-এর উপরে।",
    "AUROC interval zero cross করলেও non-inferiority pass করতে পারে, কারণ প্রশ্ন superiority নয়; unacceptable loss -0.015-এর নিচে গেছে কি না। Lower bound -0.00212, তাই worst plausible loss margin-এর তুলনায় অনেক ছোট।",
    ["AUPRC difference -0.00926 CI [-0.01160,-0.00691]", "Brier +0.00211 CI [0.00177,0.00245]", "Eight SOTA non-inferiority tests passed after Holm", "Non-inferiority != superiority"],
    ["Effect size before p-value", "Confidence interval interprets direction and uncertainty", "Adverse endpoints stay visible"],
    [
        ("CI zero cross করেও result significant কীভাবে?", "Superiority test-এ zero relevant; non-inferiority test-এ -0.015 margin relevant। Entire interval margin-এর উপরে, তাই unacceptable degradation reject করা যায়।"),
        ("Holm correction কেন?", "একই family-তে multiple test করলে chance false positive বাড়ে। Holm adjusted significance family-wise error control করে।"),
        ("AUPRC result কি primary success invalidate করে?", "Registered co-primary gates invalidate করে না, কিন্তু universal utility superiority নিষিদ্ধ করে। Clinical interpretation-এ AUPRC decline গুরুত্বপূর্ণ adverse evidence।"),
    ],
)

section_lesson(
    "4.3.1 Visualization of Results",
    "Utility-communication plot-এ CAABC-FL 56 শতাংশ upload reduction সহ AUROC dense encrypted cluster-এর কাছাকাছি। Poisoning heatmap weighted mean-এর model replacement/sign flip vulnerability দেখায়, আর robust plaintext methods near-zero delta region-এ। Membership plot DP methods chance line-এর কাছে। Scalability plot client বাড়লে upload ও latency প্রায় linear বাড়ে।",
    "Graph পড়ার নিয়ম: axis কী, lower/higher কোনটি ভালো, error bar কী, baseline কোথায়। এক chart দিয়ে সব claim করবে না। Utility plot privacy বলে না; MIA plot formal DP বলে না; poisoning plot deployment security complete বলে না।",
    ["Error bar = seed uncertainty", "Dashed 0.5 MIA chance line", "Poisoning delta zero best", "Communication x-axis actual encrypted reduction"],
    ["Visual evidence and table value consistent হতে হবে।"],
    [
        ("Trade-off graph-এ ideal point কোথায়?", "High AUROC এবং high encrypted upload reduction - plot convention অনুযায়ী upper-right। CAABC right side-এ substantial saving এবং similar AUROC দেয়।"),
        ("Poisoning heatmap-এ red কী?", "More negative AUROC delta, অর্থাৎ attack-এর কারণে clean-task discrimination বড়ভাবে কমেছে। Colorbar interpretation আগে বলতে হবে।"),
    ],
)

section_lesson(
    "4.3.2 User Evaluation and System Case Studies",
    "Formal clinician user study করা হয়নি। Qualitative system case exercise হয়েছে: authorized inference released; replay rejected; wrong purpose rejected; insufficient quorum rejected; revocation-এর পরে denied; audit chain valid; raw patient pseudonym audit-এ absent। Forged aggregate authority reject করেছে; one witness outage দুই-of-three policy tolerate করেছে; restarted witness reconcile করেছে।",
    "এগুলো usability score নয়, policy and failure-case validation। ‘User evaluation’ template heading থাকলেও thesis honestভাবে বলে human participant study নেই।",
    ["Authorized release", "Replay protection", "Purpose and consent binding", "Witness recovery", "Forged aggregate rejection"],
    ["Do not call case tests clinical user validation"],
    [
        ("User satisfaction measure কোথায়?", "করা হয়নি। Section system case studies report করে; clinician or patient usability study limitation ও future work।"),
        ("Audit-এ patient privacy কীভাবে?", "Raw patient pseudonym রাখা হয়নি; purpose, model/schema binding ও minimized event evidence hashed/signed form-এ রাখা হয়েছে।"),
    ],
)

section_lesson(
    "4.4.1 Benchmark Comparison",
    "Five locked seed same-protocol SOTA comparison-এ DP-FedAvg-CKKS, CAABC-FL, Fed-EF-SGD, Fed-EF-AMS এবং adapted SoteriaFL-SGD ছিল। AUROC 0.727 থেকে 0.732। CAABC-FL 0.730 [0.727,0.734], AUPRC 0.384 [0.376,0.393] এবং 56 শতাংশ encrypted saving। Fed-EF-SGD similar AUROC কিন্তু dense CKKS re-encoding-এর পরে saving 0।",
    "Same-protocol comparison contribution isolate করে। SoteriaFL official recurrence checked হলেও record-level DP, study weighting, delta surrogate এবং dense CKKS-এর সঙ্গে adapted; তাই original paper score replication নয়।",
    ["Eight method-metric non-inferiority tests passed after Holm", "Mixed metric direction", "CAABC only substantial actual encrypted saving in aligned SOTA set"],
    ["Baseline implementation provenance", "Adaptation disclosure", "No headline-score ranking"],
    [
        ("CAABC কি Fed-EF-SGD থেকে better?", "AUROC difference -0.0006 with CI crossing zero; AUPRC 0.0106 lower, কিন্তু actual encrypted upload 56 শতাংশ কম। তাই better overall নয়; different trade-off।"),
        ("Fed-EF-AMS-এর বিরুদ্ধে?", "CAABC AUROC প্রায় একই, AUPRC +0.0438 higher, এবং encrypted saving substantial। তবুও claim evaluated protocol-এ সীমিত।"),
        ("SoteriaFL-এর বিরুদ্ধে?", "CAABC AUROC +0.0034 interval crosses zero, AUPRC -0.0076, কিন্তু communication lower। Again mixed trade-off, universal superiority নয়।"),
    ],
)

section_lesson(
    "4.4.2 Correlation with Literature Review",
    "Results literature-এর তিনটি expectation confirm করে: DP utility cost আনে কিন্তু privacy attack signal কমায়; CKKS confidentiality দেয় কিন্তু ciphertext expansion বড়; plaintext robust aggregation strong হলেও hidden-update setting-এ compatibility সীমিত। CAABC logical-to-wire gap address করে।",
    "Literature correlation মানে number match করা নয়। Mechanism-level pattern compare করা: error feedback omitted signal conserve করে, shared block sparse transport ciphertext count কমায়, DP attack efficacy কমায়, encrypted average poisoning solve করে না।",
    ["Consistent mechanism, different dataset", "Contextual studies not direct rank"],
    ["Evidence supports integration gap"],
    [
        ("Published SOTA result-এর চেয়ে আপনার AUROC কম কেন?", "Dataset, endpoint, split, privacy budget, model এবং protocol ভিন্ন। Cross-paper raw score compare invalid। Same-protocol baselinesই direct evidence।"),
    ],
)

section_lesson(
    "4.4.3 Discussion of Improvements",
    "Measured improvement encrypted communication। Ablation-এ registered k=2 block budget, refresh 5, error feedback AUROC প্রায় 0.7303 ও 56 শতাংশ saving। One block প্রায় 76 শতাংশ saving কিন্তু utility variation বাড়ায়; four/eight block benefit কমায়; every-round refresh saving প্রায় 40 শতাংশ; refresh 10 প্রায় 58 শতাংশ; error feedback বাদে AUROC mean প্রায় 0.001 কম।",
    "Ablation দেখায় result শুধু arbitrary code path নয়। Budget, refresh এবং residual mechanism trade-off control করে। Registered configuration maximum saving নয়, balanced operating point।",
    ["More compression can hurt utility", "Frequent adaptation adds traffic", "Error feedback small but directionally useful"],
    ["Ablation supports causal design interpretation"],
    [
        ("কেন k=1 নেননি যখন 76 শতাংশ saving?", "Primary design utility stability ও communication jointly optimise করেছে। k=1 task-dependent variability বাড়ায়; k=2 registered configuration utility gate pass করে substantial saving দেয়।"),
        ("Refresh 10 বেশি saving, সেটি নেননি কেন?", "Less frequent refresh stale support risk বাড়ায়। Refresh 5 responsiveness ও overhead-এর compromise হিসেবে pre-registered।"),
    ],
)

section_lesson(
    "4.5 Discussion of Findings",
    "Objective 1 research-system scope-এ achieved: local custody, DP, CKKS, authentication, verification, registry ও witness audit operational। Objective 2 achieved with 55.998 শতাংশ actual wire reduction and AUROC non-inferiority। Objective 3 addressed with multi-seed utility, privacy, poisoning, SOTA এবং service evaluation। Evidence একই সঙ্গে limitation দেখায়: AUPRC loss, poisoning vulnerability, centralized authority, single-host service এবং no external clinical validation।",
    "Best overall interpretation: framework secure collaboration-এর several properties integrate করেছে এবং CAABC meaningful efficiency point দিয়েছে; কিন্তু security completeness বা clinical superiority প্রমাণ করেনি। Honest limitation paper-এর credibility বাড়ায়।",
    ["Contribution claim narrow and strong", "Negative evidence defines next research", "No universal SOTA"],
    ["Research system, not production clinical system"],
    [
        ("আপনার hypothesis supported?", "Joint communication-utility proposition supported। Integrated feasibility research testbed-এ supported। Universal security proposition intentionally করা হয়নি; attack evaluation strengths ও failures দুটো দেখিয়েছে।"),
        ("সবচেয়ে surprising finding?", "Logical sparse Top-k encrypted wire কমায়নি, এবং strong plaintext defence privacy-compatible ছিল না। এছাড়া CAABC AUROC preserve করলেও AUPRC ও calibration adverse হয়েছে।"),
        ("Result clinically meaningful?", "Communication and discrimination result technically meaningful; clinical effectiveness claim-এর জন্য external calibration, subgroup fairness, prospective silent trial এবং workflow evaluation দরকার।"),
    ],
)

h2("System Performance এবং Cryptographic Correctness")
p("Scientific host Windows 11 AMD64 CPU environment। Real mTLS/CKKS service-এ 16 client-এর total upload 25.28 MiB এবং sequential single-host end-to-end latency 7.080 second। Two থেকে sixteen client পর্যন্ত descriptive fit প্রায় 413.5 ms per additional client, R squared 0.9996। এটি WAN বা concurrent throughput estimate নয়। Separate OpenFHE 3-of-3 multiparty CKKS maximum error 1.34e-13 এবং missing share-এ fail-safe behaviour দেখিয়েছে।")
qa("Latency production-ready কি?", "সিদ্ধান্ত দেওয়া যায় না। Measurement single physical host, sequential submission এবং controlled vector-এ; WAN, concurrency, hardware isolation ও sustained load পরীক্ষা হয়নি।")
qa("CKKS error model utility-তে প্রভাব ফেলেছে?", "Dense DP-FedAvg ও dense DP-FedAvg-CKKS প্রায় aligned, এবং cryptographic correctness error খুব ছোট। Observed main utility loss DP clipping/noise থেকে।")

h2("Chapter 4 Rapid Revision")
key_box("Defense-এর আগে মুখস্থ রাখবে", [
    "Dense encrypted reference: AUROC 0.731, AUPRC 0.394, upload 3.314 MB/round।",
    "CAABC-FL: AUROC 0.731 core / 0.730 SOTA set, AUPRC 0.384, upload 1.458 MB/round।",
    "Actual wire reduction: 55.998%, 95% CI [55.994%, 56.002%]।",
    "AUROC difference: +0.00017, CI [-0.00212,+0.00247], margin -0.015; pass।",
    "AUPRC difference: -0.00926; adverse and statistically clear।",
    "MIA AUC CAABC 0.5035; RMIA 0.5063; LiRA 0.4963।",
    "Reconstruction RMSE plaintext 0.6844; DP 2.5467; DP+CAABC 2.7632।",
    "Model replacement vs weighted mean: AUROC delta -0.393, ASR about +0.888।",
    "Longitudinal governance: rollback 0.60, quarantine recall 1.00, final ASR change +0.0046।",
    "16 clients: 25.28 MiB upload, 7.080 s single-host latency।",
])


# ---------------------------------------------------------------------------
# CHAPTER 5
# ---------------------------------------------------------------------------
chapter_break()
h1("Chapter 5 Conclusion")
p("Conclusion নতুন result যোগ করে না। এটি research question-এর direct answer, objective achievement, contribution, limitation এবং future direction একসঙ্গে দেয়। Defense-এর শেষ impression এই chapter নির্ধারণ করে।")

section_lesson(
    "5.1 Summary of the Study",
    "Study প্রশ্ন করেছে: বহু healthcare institution source record না সরিয়ে কীভাবে jointly learn করবে এবং একই সঙ্গে patient privacy, update confidentiality, encrypted communication efficiency ও accountable model release বজায় রাখবে। Literature gap থেকে layered system এবং CAABC-FL তৈরি; full natural cohort ও adversarial tests-এ evaluate করা হয়েছে।",
    "Summary-এর logical order: why collaboration, why FL insufficient, what built, how tested, what found। এখানে সব metric repeat নয়; decisive evidence ব্যবহার করবে।",
    ["Problem and response aligned", "One paragraph version প্রস্তুত রাখবে"],
    ["Local custody", "DP", "CKKS", "CAABC", "Governed release"],
    [
        ("এক মিনিটে thesis summarize করুন।", "আমরা multi-hospital learning-এর জন্য patient-level DP, CKKS aggregation, authenticated verification এবং governed release integrate করেছি। CAABC-FL CKKS-aligned shared block compression দিয়ে dense encrypted baseline-এর তুলনায় 55.998 শতাংশ upload কমিয়েছে এবং AUROC -0.015 margin-এর মধ্যে non-inferior ছিল। Privacy attacks weakened হলেও strong poisoning encrypted mean-কে ক্ষতি করেছে; তাই result communication-utility contribution, universal security claim নয়।"),
    ],
)

section_lesson(
    "5.2.1 Key Findings",
    "Key findings: real encrypted upload reduced; joint AUROC and communication gates passed; AUPRC and calibration adverse; CKKS extra predictive loss small; logical sparsity alone ineffective; DP registered privacy attacks weakened; encryption poisoning solve করেনি; governance longitudinal attack সীমিত করলেও rollback cost ছিল; system contracts forged aggregate ও witness failure cases-এ expected behaviour দেখিয়েছে।",
    "Strong conclusion balanced। Positive evidence শুধু contribution নয়; limitation-ও key finding। Board adverse result বললে defensive না হয়ে explain করবে কেন multi-metric interpretation দরকার।",
    ["Real bytes, not estimated sparsity", "Attack-specific, not universal", "Governance has availability cost"],
    ["Nine key findings trace to evidence"],
    [
        ("সবচেয়ে defensible conclusion কী?", "Evaluated protocol-এ CAABC-FL dense protected reference-এর AUROC pre-specified margin-এর মধ্যে রেখে real CKKS upload প্রায় 56 শতাংশ কমিয়েছে।"),
        ("সবচেয়ে গুরুত্বপূর্ণ negative conclusion?", "Encrypted mean poisoning-resistant নয় এবং CAABC সব predictive metric preserve করেনি।"),
    ],
)

section_lesson(
    "5.2.2 Research Contributions",
    "Theoretical contribution privacy, confidentiality, efficiency ও integrity-কে interacting properties হিসেবে frame করে। Algorithmic contribution CAABC-FL। Practical contribution multi-service executable system। Methodological contribution paired frozen protocol, actual ciphertext accounting এবং broad attack matrix। Empirical contribution same-protocol evidence and honest trade-off।",
    "Contribution answer evidenceসহ দেবে: কী নতুন, কেন useful, কোথায় measured। ‘আমরা DP, CKKS, blockchain ব্যবহার করেছি’ contribution নয়; integration rationale ও CAABC co-design contribution।",
    ["Theoretical", "Algorithmic", "Practical", "Methodological", "Empirical"],
    ["Each contribution has explicit boundary"],
    [
        ("Contribution ও feature-এর পার্থক্য কী?", "Feature system-এ থাকা component। Contribution হলো নতুন insight বা method এবং evidence যে এটি research gap address করে। CKKS একটি feature; ciphertext-aligned adaptive common-block compression plus real-byte evidence contribution।"),
        ("Bachelor thesis হিসেবে scope কি অতিরিক্ত?", "Scope modularভাবে implement ও test করা হয়েছে। Core contribution narrow; additional security/system experiments claim boundary strengthen করেছে। Production completeness দাবি করা হয়নি।"),
    ],
)

section_lesson(
    "5.3 Limitations",
    "Limitations নয়টি major group: one sepsis cohort and two hospitals; no external or prospective validation; scientific and hardened service profiles fully unified নয়; single-host systems measurement; main service centralized decryption; no real TDX/SNP; bounded attack families and five costly seeds; custom RNG not independently audited; no clinician/patient governance study।",
    "Limitation স্বীকার মানে work invalid নয়। এটি evidence কোথায় apply করে তা বলে। Board limitation জানতে চাইলে সবচেয়ে consequential তিনটি আগে বলবে: external clinical validity, threshold key custody, encrypted poisoning robustness।",
    ["Internal validity stronger than external validity", "Research system not certified", "Formal DP under assumptions; implementation needs audit"],
    ["Dataset", "Deployment", "Cryptography", "Adversary", "Human factors"],
    [
        ("সবচেয়ে বড় limitation?", "একটি retrospective cohort এবং main service-এ centralized decryption authority। প্রথমটি clinical generalization সীমিত করে, দ্বিতীয়টি trust decentralization সীমিত করে।"),
        ("এই limitation result invalidate করে?", "না। Measured claim ওই cohort ও testbed-এর জন্য valid; শুধু broader population বা production security-তে extrapolate করা যাবে না।"),
        ("Custom RNG risk কী?", "OS-seeded AES-256-CTR path recorded ও exercised হলেও independent cryptographic audit হয়নি। Formal DP implementation assurance-এর জন্য external review দরকার।"),
    ],
)

section_lesson(
    "5.4 Future Works",
    "Priority future work: multiple external clinical cohorts ও prospective silent validation; networked dropout-tolerant t-of-n key custodians; real TDX/SNP attestation-bound release; encrypted or verifiable robust aggregation; larger-model adaptive compression; independent multi-host WAN and long-duration testing; human-centred governance and fairness analysis।",
    "Future work generic ‘more data’ নয়। প্রতিটি limitation-এর direct response। Priority order হবে external validity, threshold trust, confidential robustness, then physical production evaluation।",
    ["Networked threshold integration", "Secure robust statistics", "External calibration and fairness", "Long-duration failure recovery"],
    ["Future work must be testable"],
    [
        ("একটি future work বেছে নিলে কোনটি?", "Network-integrated dropout-tolerant threshold decryption with attested custodians, কারণ current authority trust সবচেয়ে বড় security boundary। Clinical perspective-এ parallel priority external validation।"),
        ("Poisoning future work কী?", "MPC/HE-based robust sufficient statistics, verifiable aggregation বা zero-knowledge correctness evidence, identical attack protocol-এ validation/rollback-এর সঙ্গে compare।"),
    ],
)

section_lesson(
    "5.5 Recommendations",
    "Developers actual serialized encrypted traffic report করবে, privacy-accounting ও transport confidentiality আলাদা property হিসেবে document করবে, multi-metric clinical evaluation করবে এবং defence-এর plaintext/trusted-root assumptions বলবে। Consortium identity, release quorum, revocation, incident response এবং audit retention আগে define করবে। Regulators local data retention-কে complete privacy claim হিসেবে নেবে না।",
    "Recommendation তিন stakeholder-এর জন্য: engineers, institutions, reviewers/regulators। Evidence ছাড়া blockchain বা AI safety label যথেষ্ট নয়।",
    ["Technical reporting", "Institutional governance", "Policy oversight", "Reproducible research"],
    ["Operational policy before training"],
    [
        ("Hospital consortium-এর প্রথম কাজ কী?", "Participant identity, legal purpose, data and model ownership, decryption authority, release quorum, revocation এবং incident-response policy লিখিতভাবে agree করা।"),
    ],
)

section_lesson(
    "5.6 Concluding Remarks",
    "Trustworthy collaborative intelligence একটি technique থেকে আসে না। Algorithms, privacy accounting, cryptography, systems engineering এবং institutional responsibility evidence-aligned হতে হয়। CAABC-FL একটি concrete step: real encrypted traffic কমায় এবং bounded utility preserve করে; stronger threshold governance ও confidential robustness এখনও প্রয়োজন।",
    "শেষ বক্তব্য confident কিন্তু honest হবে। ‘আমরা problem solve করেছি’ না বলে ‘আমরা evaluated setting-এ একটি defensible operating point demonstrate করেছি’ বলবে।",
    ["End with contribution plus boundary plus next step"],
    ["Avoid overclaiming"],
    [
        ("Thesis-এর final takeaway কী?", "Privacy-preserving multi-institutional learning শুধু data local রাখার বিষয় নয়; privacy, encrypted efficiency, poisoning risk এবং accountable governance একসঙ্গে design ও measure করতে হয়।"),
    ],
)

h2("Chapter 5 Rapid Revision")
key_box("Conclusion-এর পাঁচটি বাক্য", [
    "Integrated research system secure collaboration-এর major layers একত্র করেছে।",
    "CAABC-FL real CKKS upload 55.998 শতাংশ কমিয়েছে।",
    "AUROC non-inferiority passed, কিন্তু AUPRC ও calibration adverse হয়েছে।",
    "DP registered privacy attacks weakened করেছে, কিন্তু encryption poisoning solve করেনি।",
    "External clinical validation, networked threshold custody এবং encrypted robust aggregation next priorities।",
])


# ---------------------------------------------------------------------------
# FINAL REVISION
# ---------------------------------------------------------------------------
chapter_break()
h1("Final Rapid Revision এবং Mock Defense")

h2("A. এক পাতার Key Numbers")
two_col_table(
    ["বিষয়", "সংখ্যা", "কী বোঝায়"],
    [
        ("Dataset", "40,336 patients, 2 natural hospitals", "Full retrospective cohort"),
        ("Features", "280 from 24-hour history", "Matched patient representation"),
        ("Model", "280-64-1, 18,049 parameters", "Compact MLP"),
        ("Training", "20 rounds, 1 local epoch", "Frozen protocol"),
        ("DP", "ε=1.935, δ=10^-5", "Patient-level accountant"),
        ("CKKS", "8192, 60-40-40-60, scale 2^40", "Packed approximate encryption"),
        ("CAABC", "B=256, c=2, refresh=5", "Aligned adaptive blocks"),
        ("Dense upload", "3.314 MB/round", "Protected reference"),
        ("CAABC upload", "1.458 MB/round", "Adaptive encrypted upload"),
        ("Wire reduction", "55.998%", "Co-primary gate passed"),
        ("AUROC difference", "+0.00017 [-0.00212,+0.00247]", "Non-inferior to -0.015 margin"),
        ("AUPRC difference", "-0.00926", "Adverse trade-off"),
        ("CAABC MIA", "0.5035", "Registered attack near chance"),
        ("Reconstruction", "0.684 -> 2.763 RMSE", "DP plus CAABC harder to reconstruct"),
        ("Worst poisoning", "AUROC -0.393; ASR +0.888", "Encrypted mean vulnerable"),
        ("Longitudinal gate", "rollback 0.60; quarantine recall 1.00", "Control with availability cost"),
        ("Service 16 clients", "25.28 MiB; 7.080 s", "Single-host sequential diagnostic"),
    ],
    [1.55, 2.2, 2.8],
)

h2("B. Contribution Evidence Map")
two_col_table(
    ["Contribution", "Evidence", "সীমা"],
    [
        ("CAABC-FL algorithm", "Shared CKKS-aligned blocks, adaptive refresh, error feedback", "One model and cohort"),
        ("Real communication saving", "55.998% serialized ciphertext reduction", "Current CKKS parameters and model"),
        ("Utility retention", "AUROC NI gate passed", "AUPRC and F1 declined"),
        ("Privacy evaluation", "MIA/LiRA/RMIA plus 300-step reconstruction", "Attack-specific; DP formal guarantee separate"),
        ("Poisoning evaluation", "870 one-round conditions plus longitudinal backdoor", "No formal Byzantine proof"),
        ("Trust architecture", "Signatures, mTLS, recomputation, registry, witnesses", "Authority centrally keyed"),
        ("Threshold evidence", "OpenFHE native 3-of-3, missing share fails", "Not service-integrated or dropout-tolerant"),
    ],
    [1.7, 2.8, 2.05],
)

h2("C. Board-এর 20টি High-Probability প্রশ্ন")
final_qas = [
    ("CAABC-FL full form কী?", "Ciphertext-Aligned Adaptive Block Compression for Federated Learning।"),
    ("Title-এর সঙ্গে কাজ align করে?", "হ্যাঁ। Hybrid কারণ DP, CKKS, compression, trust and audit combine; privacy-preserving কারণ patient-level DP and encrypted aggregation; secure collaboration কারণ authentication, verification and governed release; multi-institutional কারণ natural hospital silos jointly train।"),
    ("Raw patient data কোথায় যায়?", "Hospital boundary ছাড়ে না; only privatized, compressed, encrypted model update and signed metadata যায়।"),
    ("Main baseline কেন dense DP-FedAvg-CKKS?", "এটি proposed method-এর একই patient-level DP and CKKS protection রাখে; difference প্রধানত CAABC compression, তাই causal comparison fair।"),
    ("Privacy budget কী?", "Frozen patient-level accountant-এ ε=1.935, δ=10^-5, clipping 1.0 এবং noise multiplier 1.2।"),
    ("Why not only HE?", "HE server থেকে update গোপন করে, কিন্তু final model leakage বা patient influence bound করে না; DP সেই guarantee দেয়।"),
    ("Why not only DP?", "DP formal influence bound দেয়, কিন্তু server individual noisy update দেখতে পারে; CKKS institutional update confidentiality দেয়।"),
    ("CAABC score plaintext কি?", "Block importance already privatized update থেকে derive এবং aggregate করা হয়; unclipped raw gradient inspect করা হয় না।"),
    ("Non-inferiority কী?", "Proposed method reference-এর চেয়ে acceptable margin-এর বেশি খারাপ নয় তা test করা। আমাদের AUROC lower CI -0.015 margin-এর উপরে।"),
    ("p-value বলেননি কেন?", "Primary decision pre-specified confidence-bound gates-এর উপর; effect size and interval practical magnitude দেখায়। p-value alone যথেষ্ট নয়।"),
    ("AUPRC কমা কীভাবে defend করবেন?", "লুকাব না। এটি positive-class trade-off; তাই claim AUROC-preserving communication, overall clinical superiority নয়।"),
    ("Encryption poisoning আটকায়?", "না। It hides content, not maliciousness; weighted encrypted mean strong replacement attack-এ vulnerable।"),
    ("FLTrust direct baseline নয় কেন?", "Trusted labelled root data and per-client update geometry লাগে, sum-only encrypted server-এর assumptions-এর সঙ্গে conflict।"),
    ("MIA AUC 0.5 কী বোঝায়?", "Registered attacker chance-level discrimination করেছে; formal privacy proof নয়।"),
    ("Threshold system আছে?", "Separate 3-of-3 OpenFHE demonstration আছে; main TenSEAL service centrally keyed and not dropout tolerant।"),
    ("Blockchain ব্যবহার করেছেন?", "No public blockchain; signed permissioned witness ledger with two-of-three receipts।"),
    ("Production ready?", "No. Research system; external clinical validation, multi-host, attestation, audit and threshold integration বাকি।"),
    ("সবচেয়ে বড় contribution?", "Ciphertext-aware adaptive compression that turns sparse logic into measured encrypted-wire saving without unacceptable AUROC loss।"),
    ("সবচেয়ে বড় limitation?", "Single retrospective cohort and centrally held main-service decryption key।"),
    ("পরের experiment?", "External multi-hospital validation এবং networked dropout-tolerant threshold custodians with encrypted robust aggregation।"),
]
for qtext, atext in final_qas:
    qa(qtext, atext)

h2("D. উত্তর দেওয়ার Safe Language")
two_col_table(
    ["Avoid", "Use instead"],
    [
        ("Our system is fully secure", "Under the registered threat model, the evaluated controls reduced specific risks"),
        ("We beat the SOTA", "We passed same-protocol non-inferiority tests and achieved lower encrypted upload"),
        ("DP stopped all attacks", "The formal DP accountant bounded patient influence and registered attacks became less effective"),
        ("Threshold decryption is deployed", "A separate native 3-of-3 OpenFHE experiment demonstrated multiparty decryption"),
        ("The system is production ready", "The executable research system validates the architecture; production validation remains future work"),
        ("Compression improved privacy", "Compression followed the DP mechanism; formal privacy comes from the accountant"),
    ],
    [3.0, 3.6],
)

h2("E. শেষ দিনের Preparation Plan")
number("Key Numbers table তিনবার মুখে বলো, unit সহ।")
number("Architecture flow diagram না দেখে actor ও data object অনুসরণ করে ব্যাখ্যা করো।")
number("AUROC non-inferiority এবং AUPRC adverse result পাশাপাশি explain করো।")
number("DP বনাম CKKS বনাম authentication বনাম audit - চারটির কাজ এক বাক্যে আলাদা করো।")
number("Poisoning-proof, fully decentralized, production ready এবং universal SOTA - এই চারটি claim কেন করবে না, অনুশীলন করো।")
number("প্রত্যেকে নিজের section ছাড়াও Abstract, contribution, limitation এবং future work উত্তর দিতে পারবে।")

h2("Final Closing Statement for Defense")
closing = p("Our evidence shows that privacy-preserving federated learning must be judged jointly by utility, privacy, real ciphertext traffic, and integrity. Under the evaluated protocol, CAABC-FL reduced serialized CKKS upload by about 56 percent while remaining AUROC-non-inferior to dense protected training. This is a bounded research contribution, not a claim of universal SOTA or production readiness; external validation, networked threshold custody, and confidential robust aggregation remain future work.", italic=True)
closing.paragraph_format.keep_together = True
closing.paragraph_format.space_after = Pt(0)
for run in closing.runs:
    set_run_font(run, 9.2, italic=True)


# Final formatting pass
for para in doc.paragraphs:
    para.paragraph_format.widow_control = True
    for run in para.runs:
        if run.font.name is None:
            set_run_font(run)

for table in doc.tables:
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.widow_control = True
                for run in para.runs:
                    if run.font.name is None:
                        set_run_font(run)

doc.core_properties.title = "CAABC-FL Thesis Defense Preparation Guide"
doc.core_properties.subject = "Section-wise Bangla lessons and defense questions from Abstract to Conclusion"
doc.core_properties.author = "CAABC-FL Thesis Team"
doc.core_properties.keywords = "CAABC-FL, federated learning, differential privacy, CKKS, defense preparation"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
